"""Pure provisional reporting over immutable parsed amounts.

build_report returns months (cash/business aggregates), cash_rows, profit_rows,
issues, metrics, details, cumulative, source_inventory, templates and a separately
labelled ledger_summary. Decisions may confirm category, period, payer, exclusion
and evidence links. Original amounts/dates remain immutable; reviewed image
discount allocations and month-end accrual dates are separate derived values. Final profit
is unavailable while blocking source, classification, period or overlap issues
remain. Owner distributions are separate from operating expenses and deducted
only in the explicitly labelled template result. Export is an in-memory XLSX.
"""
from __future__ import annotations

import io
import re
import calendar
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation

from .image_facts import effective_image_fact

ZERO = Decimal(0)
NON_OPERATING = {"capital", "capital_return", "loan", "loan_repayment", "internal_transfer"}
EXPENSES = {"payroll", "social_insurance", "bank_fee", "tax", "laundry_supplier",
            "cleaning_supplier", "property", "maintenance", "utility", "water",
            "electricity", "supplies", "rent", "operating_expense", "cleaning_supplier_cost", "laundry_supplier_cost"}
INCOME = {"operating_income", "revenue", "ota"}
PERIOD = re.compile(r"^\d{4}-(?:0[1-9]|1[0-2])$")


def _dec(value):
    try:
        amount = Decimal(str(value))
        return amount if amount.is_finite() else None
    except InvalidOperation:
        return None


def _fmt(value):
    return format(value.quantize(Decimal("0.01")), "f")


def _month(value):
    return value if isinstance(value, str) and PERIOD.fullmatch(value) else None


def _category_family(value):
    return "utility" if value in {"water", "electricity", "utility"} else {"cleaning_supplier_cost": "cleaning_supplier", "laundry_supplier_cost": "laundry_supplier"}.get(value, value)


def _safe_description(fact, category, month):
    if fact.get("kind") != "bank_transaction":
        return str(fact.get("description") or category)[:300]
    # Raw bank narratives can contain names/accounts. Preserve only a business
    # reference deliberately extracted by the source parser, plus classification.
    ref = str(fact.get("source_reference") or "")
    ref = ref if re.fullmatch(r"[A-Za-z0-9-]{6,80}", ref) else ""
    return " / ".join(v for v in [category, month, f"业务单号 {ref}" if ref else None] if v)


def build_report(sources: list[dict], decisions: dict, months: list[str], ledger_summary: dict | None = None) -> dict:
    report = dict(months=[], cash_rows=[], profit_rows=[], issues=[], metrics={}, details=[],
                  source_inventory=[], source_summaries=[], templates=[], ledger_summary={"basis": "system_service_fees_not_supplier_actual_costs", "data": ledger_summary},
                  notes=["现金按银行交易日统计；业务贡献按已确认的业务月份统计。",
                         "OTA及供应方明细与银行款项为可能重叠的证据，未确认对应关系时不重复相加。",
                         "公司自身费用与代业主支付费用分列；业主分成净额须另扣代付，分成毛额已含代付时不重复扣减。",
                         "原利润模板只作对照，不构成已确认会计结果。"])
    requested = sorted({_month(m) for m in months if _month(m)})
    facts = {}

    def issue(code, message, row=None, *, blocking=True, source_id=None):
        report["issues"].append(dict(code=code, message=message, blocking=blocking,
            severity="error" if blocking else "info", source_id=source_id or (row or {}).get("source_id"),
            fact_key=(row or {}).get("fact_key"), month=(row or {}).get("business_month")))

    for source in sources:
        parsed = source.get("parsed") or {}
        source_id = source.get("source_id")
        report["source_inventory"].append(dict(source_id=source_id, filename=source.get("filename"),
            kind=parsed.get("kind") or source.get("kind"), fact_count=len(parsed.get("facts", [])),
            issue_count=len(parsed.get("issues", []))))
        report["templates"].extend(dict(t, source_id=source_id) for t in parsed.get("templates", []))
        report["source_summaries"].extend(dict(s, source_id=source_id, countable=False) for s in parsed.get("summaries", []))
        image_candidates = [f for f in parsed.get("facts", []) if f.get("kind") == "image_expense_candidate"]
        for original in parsed.get("issues", []):
            code = original.get("code", "SOURCE_ISSUE")
            if image_candidates and code in {"image_requires_verification", "image_candidates_unconfirmed", "image_expense_category_unknown", "image_discount_allocation_required"}:
                continue  # Specific candidate checks below replace generic OCR caveats.
            if code == "expense_payer_unconfirmed":
                expense_facts = [f for f in parsed.get("facts", []) if f.get("kind") == "source_expense"]
                payer_values = {"company", "company_advanced", "公司", "公司垫付", "owner", "owner_paid", "业主", "业主已付"}
                if expense_facts and all((decisions.get(f.get("key")) or {}).get("status") == "confirmed"
                    and ((decisions.get(f.get("key")) or {}).get("include") is False
                         or (decisions.get(f.get("key")) or {}).get("payer") in payer_values) for f in expense_facts):
                    continue
            # Parser domain caveats are handled by explicit decisions below.
            blocking = original.get("severity") == "error" or code.startswith(("INVALID_", "MISSING_BANK_", "BOTH_BANK_", "NEGATIVE_BANK_", "BANK_RUNNING_", "BANK_BALANCE_", "UNSUPPORTED_", "ocr_", "image_"))
            issue(code, original.get("message") or code, blocking=blocking, source_id=source_id)
        for fact in parsed.get("facts", []):
            key = fact.get("key")
            if not key:
                issue("missing_fact_key", "来源明细缺少稳定证据编号", source_id=source_id)
                continue
            if key in facts:
                issue("duplicate_source_fact", "重复来源证据已按稳定编号计入一次", source_id=source_id, blocking=False)
                continue
            decision = decisions.get(key) or {}
            confirmed = decision.get("status") == "confirmed"
            fact = effective_image_fact(fact, decision)
            category = decision.get("category", fact.get("category", "unknown")) if confirmed else fact.get("category", "unknown")
            business_month = _month(decision.get("business_month", fact.get("business_month"))) if confirmed else _month(fact.get("business_month"))
            day = str(fact.get("date") or "")
            row = dict(source_id=source_id, fact_key=key, kind=fact.get("kind"), date=fact.get("date"),
                cash_month=_month(day[:7]), business_month=business_month, category=category,
                amount=fact.get("amount"), direction=fact.get("direction"), channel=fact.get("channel"),
                confirmed=confirmed, include=decision.get("include", True) if confirmed else True,
                payer=decision.get("payer") if confirmed else None,
                paid_by=decision.get("paid_by") if confirmed else None,
                independent_cost=confirmed and decision.get("independent_cost") is True,
                owner_distribution_basis=decision.get("owner_distribution_basis") if confirmed else None,
                description=_safe_description(fact, category, business_month),
                lineage=[dict(source_id=source_id, fact_key=key, sheet=fact.get("sheet"),
                              row1based=fact.get("row1based", fact.get("row")))],
                fields=fact.get("fields") or {}, decision_note=str(decision.get("note") or "")[:500])
            if fact.get("kind") == "image_expense_candidate":
                row["room_ref"] = decision.get("room_ref", fact.get("room_ref")) if confirmed else fact.get("room_ref")
                row["payment_date"] = fact.get("date")
                row["date_basis"] = decision.get("date_basis") if confirmed else None
                if confirmed and decision.get("expense_date"):
                    row["date"] = decision["expense_date"]
                elif confirmed and decision.get("date_basis") == "month_end_accrual" and business_month:
                    year, number = map(int, business_month.split("-"))
                    row["date"] = date(year, number, calendar.monthrange(year, number)[1]).isoformat()
            if fact.get("kind") == "image_text":
                row["candidate_evidence_only"] = bool(image_candidates)
                if not image_candidates and row["include"]:
                    issue("image_requires_verification", "图片尚无可核实费用候选，请补充原始金额、用途和所属月份", row)
            row["lineage"].extend(dict(ref, source_id=source_id, fact_key=key) for ref in fact.get("evidence_refs", []))
            facts[key] = (row, decision)
            report["details"].append(row)
    parent = {key: key for key in facts}

    def find(key):
        while parent[key] != key:
            parent[key] = parent[parent[key]]
            key = parent[key]
        return key

    for key, (row, decision) in facts.items():
        if not row["confirmed"]:
            continue
        links = decision.get("matching_fact_keys") or []
        if not isinstance(links, list):
            links = []
            issue("invalid_evidence_links", "对应证据编号应为列表", row)
        links = links + [decision[field] for field in ("counterpart_key", "duplicate_of") if decision.get(field)]
        if row["independent_cost"] and links:
            issue("invalid_evidence_links", "独立费用不能同时关联为同一笔费用，请先澄清对应关系", row)
            row["independent_cost"] = False
        for other in links:
            if other not in facts:
                issue("missing_linked_evidence", "已确认对应的来源证据不存在", row)
            else:
                parent[find(other)] = find(key)
        if decision.get("duplicate_of") in facts and decision.get("duplicate_of") != key:
            row["include"] = False
            issue("confirmed_duplicate_excluded", "已确认重复交易只计入对应原交易一次", row, blocking=False)

    groups = defaultdict(list)
    for key, (row, _) in facts.items():
        groups[find(key)].append(row)
        if not row["include"]:
            continue
        if row["kind"] == "bank_transaction":
            amount = _dec(row["amount"])
            if amount is None or amount < 0 or row["direction"] not in {"in", "out"}:
                issue("invalid_cash_fact", "银行金额或方向无效，未纳入现金汇总", row)
            elif not row["cash_month"]:
                issue("cash_month_unknown", "交易日期不完整，未纳入月度现金汇总", row)
            else:
                report["cash_rows"].append(row)

    selected_groups = []
    invalid_images = set()
    for key, (row, decision) in facts.items():
        if row["kind"] != "image_expense_candidate" or not row["include"]:
            continue
        def image_issue(code, message):
            issue(code, message, row)
            invalid_images.add(key)
        if not row["confirmed"] or decision.get("candidate_confirmed") is not True:
            image_issue("image_amount_unconfirmed", "图片费用金额和用途尚未逐项核实，不能计入经营费用")
            continue
        if row["paid_by"] not in {"company", "owner", "公司", "业主"}:
            image_issue("image_payment_unconfirmed", "请明确图片费用实际由公司还是业主支付")
        amount = _dec(row["amount"])
        if amount is None or amount <= ZERO:
            image_issue("business_amount_unknown", "图片费用金额应为已核实的正数，退款或零金额需单独说明")
        try:
            date.fromisoformat(str(row["date"]))
        except ValueError:
            image_issue("image_date_unconfirmed", "请核实费用日期，或明确按已确认业务月末归集；图片时间不等于付款日")
        fields = row["fields"]
        if fields.get("discount_allocation_required") and fields.get("applied_discount_allocation") != "proportional":
            image_issue("image_discount_unconfirmed", "明细尚未分摊总优惠，请确认按原金额比例分摊或选择实付汇总")
        if fields.get("allocation_required") and row["payer"] in {"owner", "业主"}:
            image_issue("image_owner_allocation_required", "汇总涉及多个房间，需先核实各房金额及优惠分摊，不能全部扣给一个业主")
        alternatives = fields.get("mutually_exclusive_with", [])
        if any(other in facts and facts[other][0]["include"] and facts[other][0]["confirmed"] and facts[other][1].get("candidate_confirmed") is True for other in alternatives):
            image_issue("image_alternative_conflict", "同一图片的汇总与明细被同时选中，请明确保留一种口径并排除另一种")
    for group in groups.values():
        included = [r for r in group if r["include"] and r["kind"] not in {"cleaning_work", "image_text", "image_expense_summary", "image_detail_total"}]
        if not included:
            continue
        bank = [r for r in included if r["kind"] == "bank_transaction"]
        source = [r for r in included if r["kind"] in {"platform_order", "platform_payout", "source_expense", "image_expense_candidate"}]
        group_blocked = False
        representations = defaultdict(list)
        for row in source:
            representations[(row["source_id"], row["kind"])].append(row)
        if len(representations) > 1:
            totals = {sum((_dec(r["amount"]) or ZERO for r in rows), ZERO) for rows in representations.values()}
            if len(totals) > 1:
                issue("counterpart_amount_mismatch", "同笔事项在不同来源中的合计不一致，请核实各表覆盖范围、优惠或拆分依据", source[0])
                group_blocked = True
        # Rows in one statement may be a real split of one payment. Separate
        # statements are alternate representations, never additive duplicates.
        canonical = sorted(representations.values(), key=lambda rows: (
            not all(r["confirmed"] for r in rows), -len(rows),
            str(rows[0]["source_id"]), str(rows[0]["kind"])))
        representative = canonical[0] if canonical else []
        # A confirmed counterpart decision covers linkage only, never silently
        # confirms the other row's expense ownership, period or category.
        selected = representative if source and all(r["confirmed"] for r in source) else bank or representative
        confirmed_months = {r["business_month"] for r in included if r["confirmed"] and r["business_month"]}
        if len(confirmed_months) > 1:
            issue("counterpart_business_month_conflict", "同笔关联事项被确认到不同业务月份，请统一归属月份；银行实际到账日期保持不变", included[0])
            group_blocked = True
        category_aliases = {"cleaning_supplier_cost": "cleaning_supplier", "laundry_supplier_cost": "laundry_supplier", "property_fee": "property"}
        categories = {category_aliases.get(r["category"], r["category"]) for r in included if r["confirmed"] and r["category"] not in {None, "unknown"}}
        compatible_categories = len(categories) <= 1 or categories <= INCOME or ("utility" in categories and categories <= {"utility", "water", "electricity"})
        if not compatible_categories:
            issue("counterpart_category_conflict", "同笔关联事项的已确认类别不一致，请核实是否同一笔业务；本金往来不能与费用合并，明确水费与电费也不能相互替代", included[0])
            group_blocked = True
        if bank and source:
            bank_total = sum((_dec(r["amount"]) or ZERO for r in bank), ZERO)
            source_total = sum((_dec(r["amount"]) or ZERO for r in representative), ZERO)
            if bank_total != source_total:
                issue("counterpart_amount_mismatch", "关联的银行款项与来源明细金额不符，需补充退款、手续费或拆分依据", selected[0])
            if len({r["direction"] for r in bank}) > 1:
                issue("mixed_counterpart_direction", "同组关联证据混有收款及付款，需明确资金分配", selected[0])
            conflicting_payments = [r for r in included if r["paid_by"] in {"owner", "业主"}]
            if conflicting_payments:
                issue("counterpart_payment_conflict", "费用已标记实际由业主付款，但关联了公司银行付款，需核实是否报销或关联错误", conflicting_payments[0])
                group_blocked = True
        if len(included) > 1:
            supplier_rows = [r for r in included if r["kind"] in {"source_expense", "image_expense_candidate"}]
            for expense in supplier_rows:
                if expense["payer"] not in {"company", "company_advanced", "公司", "公司垫付", "owner", "业主"}:
                    issue("expense_payer_unconfirmed", "关联费用的承担方尚未确认，其他来源或银行付款不能替代公司或业主承担的说明", expense)
                    group_blocked = True
            parties = {"owner" if r["payer"] in {"owner", "业主"} else "company"
                for r in included if r["payer"] in {"owner", "业主", "company", "company_advanced", "公司", "公司垫付"}}
            if supplier_rows and len(parties) > 1:
                issue("counterpart_payer_conflict", "同组来源对费用承担方的说明不一致，请核实公司或业主承担后再计入", supplier_rows[0])
                group_blocked = True
            payment_parties = {"owner" if r["paid_by"] in {"owner", "业主"} else "company"
                for r in supplier_rows if r["paid_by"] in {"owner", "业主", "company", "公司"}}
            if len(payment_parties) > 1:
                issue("counterpart_payment_conflict", "同组来源对实际付款方的说明不一致，请核实公司支付、业主支付或报销关系", supplier_rows[0])
                group_blocked = True
        if group_blocked:
            selected = []
        lineage = [ref for r in included for ref in r["lineage"]]
        selected_groups.append((included, selected, lineage))

    for included, selected, lineage in selected_groups:
        for row in selected:
            if row["fact_key"] in invalid_images:
                continue
            if not row["confirmed"]:
                issue("classification_unconfirmed", "业务分类及会计处理尚未确认", row)
                continue
            category = row["category"]
            if category in NON_OPERATING:
                continue
            if category not in EXPENSES | INCOME | {"owner_distribution"}:
                issue("category_unresolved", "未识别为可计入业务结果的明确分类", row)
                continue
            if not row["business_month"]:
                issue("business_month_unknown", "缺少业务归属月份，不能以到账月份代替", row)
                continue
            amount = _dec(row["amount"])
            if amount is None:
                issue("business_amount_unknown", "业务金额无法可靠识别", row)
                continue
            owner_advance = False
            if row["kind"] == "bank_transaction" and category in EXPENSES:
                if row["paid_by"] in {"owner", "业主"}:
                    issue("counterpart_payment_conflict", "公司银行流水与实际由业主付款的说明冲突，请核实账户归属或报销关系；原银行现金事实保留", row)
                    continue
                if row["payer"] in {"owner", "业主"}:
                    if row["paid_by"] in {"company", "公司"}:
                        owner_advance = True
                    else:
                        issue("owner_expense_payment_unknown", "业主承担费用的实际付款方未明确确认，不能直接认定公司代付", row)
                        continue
            if row["kind"] in {"source_expense", "image_expense_candidate"} and row["payer"] not in {"company", "company_advanced", "公司", "公司垫付"}:
                if row["payer"] in {"owner", "业主"}:
                    if row["paid_by"] not in {"company", "owner", "公司", "业主"}:
                        issue("owner_expense_payment_unknown", "费用承担方为业主，不代表业主已付款；实际公司垫付及业主结算扣除仍待对应", row)
                    elif row["paid_by"] in {"owner", "业主"}:
                        issue("owner_borne_expense_excluded", "业主承担费用未计入公司自身成本；已记录实际付款方，业主结算扣除仍须按关联证据核对", row, blocking=False)
                    else:
                        owner_advance = True
                    if not owner_advance:
                        continue
                else:
                    issue("expense_payer_unconfirmed", "费用承担方尚未确认；供应方收款不能证明公司承担费用或公司已垫付", row)
                    continue
            if row["kind"] != "bank_transaction" and not (category in EXPENSES and row["independent_cost"]):
                candidates = [other for group, _, _ in selected_groups for other in group
                    if (other["kind"] == "bank_transaction" or row["kind"] == "image_expense_candidate" and other["kind"] == "source_expense") and other not in included
                    and other["include"] and (other["business_month"] in {None, row["business_month"]})
                    and ((category in INCOME and other["category"] in INCOME)
                         or _category_family(category) == _category_family(other["category"]))]
                if candidates:
                    issue("possible_duplicate_business_evidence", "存在同类来源款项尚未建立对应关系；来源金额暂不叠加", row)
                    continue
            role = "owner_advance" if owner_advance else "owner_share" if category == "owner_distribution" else "revenue" if category in INCOME else "expense"
            if row["kind"] == "bank_transaction" and ((role == "revenue" and row["direction"] != "in") or (role != "revenue" and row["direction"] != "out")):
                issue("category_direction_conflict", "已确认分类与银行收付方向不一致，退款需明确处理", row)
                continue
            report["profit_rows"].append(dict(row, role=role, amount=_fmt(amount), lineage=lineage))

    if not requested:
        requested = sorted({r["cash_month"] for r in report["cash_rows"]} | {r["business_month"] for r in report["profit_rows"]})
    if not facts:
        issue("no_evidence", "尚无可核对明细证据")
    elif not report["cash_rows"]:
        issue("bank_evidence_missing", "尚无可核对银行流水；OTA结算、供应方收款或保洁记录不足以覆盖公司整体收支，不能定稿利润")
    for month in requested:
        if not any(r["cash_month"] == month for r in report["cash_rows"]) and not any(r["business_month"] == month for r in report["profit_rows"]):
            issue("requested_month_evidence_missing", f"{month} 尚无已识别现金或已确认业务事实；空白不表示该月收入和成本为零")
    owner_bases = {}
    for month in requested:
        shares = [r for r in report["profit_rows"] if r["business_month"] == month and r["role"] == "owner_share"]
        advances = [r for r in report["profit_rows"] if r["business_month"] == month and r["role"] == "owner_advance"]
        bases = {r["owner_distribution_basis"] for r in shares}
        if not shares:
            owner_bases[month] = "missing" if advances else "not_applicable"
            if advances:
                issue("owner_settlement_missing", "已确认公司代业主支付费用，但缺少当月业主分成金额及净额或毛额口径；请核实当月结算，不能把空白当作无需分成", advances[0])
        elif not bases <= {"net", "gross"}:
            owner_bases[month] = "unknown"
            issue("owner_distribution_basis_unknown", "请确认当月业主分成是已扣除代付费用的净额，还是扣费前毛额；目前按净额口径暂算，不能定稿", shares[0])
        elif len(bases) > 1:
            owner_bases[month] = "mixed"
            issue("owner_distribution_basis_mixed", "当月业主分成混有净额和毛额，无法确定代付费用是否已包含；请统一结算口径或补充逐业主的代付分配依据", shares[0])
        else:
            owner_bases[month] = next(iter(bases))
    blocking_count = sum(i["blocking"] for i in report["issues"])
    for month in requested:
        cash = [r for r in report["cash_rows"] if r["cash_month"] == month]
        profit = [r for r in report["profit_rows"] if r["business_month"] == month]
        cash_in = sum((_dec(r["amount"]) for r in cash if r["direction"] == "in"), ZERO)
        cash_out = sum((_dec(r["amount"]) for r in cash if r["direction"] == "out"), ZERO)
        roles = {role: sum((_dec(r["amount"]) for r in profit if r["role"] == role), ZERO) for role in ("revenue", "expense", "owner_share", "owner_advance")}
        categories = {}
        for category in sorted({r["category"] for r in cash}):
            amounts = {direction: sum((_dec(r["amount"]) for r in cash if r["category"] == category and r["direction"] == direction), ZERO) for direction in ("in", "out")}
            categories[category] = {k: _fmt(v) for k, v in amounts.items()}
        operating = roles["revenue"] - roles["expense"]
        advance_deduction = ZERO if owner_bases[month] == "gross" else roles["owner_advance"]
        template = None if owner_bases[month] == "mixed" else operating - roles["owner_share"] - advance_deduction
        report["months"].append(dict(month=month, cash_in=_fmt(cash_in) if cash else None, cash_out=_fmt(cash_out) if cash else None,
            net_cash=_fmt(cash_in - cash_out) if cash else None, cash_categories=categories,
            confirmed_revenue=_fmt(roles["revenue"]), confirmed_expenses=_fmt(roles["expense"]),
            confirmed_owner_share=_fmt(roles["owner_share"]), confirmed_owner_advances=_fmt(roles["owner_advance"]),
            owner_distribution_basis=owner_bases[month], confirmed_operating_result=_fmt(operating),
            provisional_template_result=_fmt(template) if template is not None else None, final_profit=None if blocking_count or template is None else _fmt(template),
            cash_fact_keys=[r["fact_key"] for r in cash], profit_fact_keys=[r["fact_key"] for r in profit]))
    sums = ("cash_in", "cash_out", "net_cash", "confirmed_revenue", "confirmed_expenses", "confirmed_owner_share", "confirmed_owner_advances", "confirmed_operating_result")
    report["cumulative"] = {key: None if not report["months"] or any(m[key] is None for m in report["months"])
        else _fmt(sum((Decimal(m[key]) for m in report["months"]), ZERO)) for key in sums}
    report["cumulative"]["provisional_template_result"] = None if any(m["provisional_template_result"] is None for m in report["months"]) else _fmt(sum((Decimal(m["provisional_template_result"]) for m in report["months"]), ZERO))
    report["cumulative"]["final_profit"] = None if blocking_count else report["cumulative"]["provisional_template_result"]
    report["status"] = "provisional" if blocking_count else "ready_for_review"
    report["metrics"] = dict(source_count=len(sources), fact_count=len(facts), cash_row_count=len(report["cash_rows"]),
        profit_row_count=len(report["profit_rows"]), open_issue_count=blocking_count,
        confirmed_count=sum(r["confirmed"] for r, _ in facts.values()),
        excluded_nonoperating_count=sum(r["confirmed"] and r["category"] in NON_OPERATING for r, _ in facts.values()))
    missing_cash_months = [m["month"] for m in report["months"] if not m["cash_fact_keys"]]
    cash_pending = "部分月份缺少银行流水" if any(m["cash_fact_keys"] for m in report["months"]) else "缺少银行流水"
    def cash_value(value):
        return value + " 元" if value is not None else cash_pending
    report["chat_metrics"] = [
        dict(label="所选月份银行收款", value=cash_value(report["cumulative"]["cash_in"])),
        dict(label="所选月份银行付款", value=cash_value(report["cumulative"]["cash_out"])),
        dict(label="所选月份净现金", value=cash_value(report["cumulative"]["net_cash"])),
        dict(label="待确认问题", value=str(blocking_count)),
        dict(label="利润状态", value="暂不能定稿" if blocking_count else "可复核"),
    ]
    report["chat_details"] = [dict(label=m["month"] + " 核对结果", value=
        (f"已提供流水中的银行收款 {m['cash_in']} 元、付款 {m['cash_out']} 元、净现金 {m['net_cash']} 元；" if m['cash_fact_keys'] else "缺少该月银行流水，收款、付款和净现金待核实；") +
        f"已确认业务收入 {m['confirmed_revenue']} 元、公司自身费用 {m['confirmed_expenses']} 元、代业主支付 {m['confirmed_owner_advances']} 元、业主分成 {m['confirmed_owner_share']} 元。"
        + ("业主分成混合口径尚未核实，暂不能计算利润。" if m["provisional_template_result"] is None
           else f"已确认部分暂算 {m['provisional_template_result']} 元；完整利润待确认。" if m["profit_fact_keys"] and blocking_count
           else "完整利润待确认，未确认部分不能按零处理。" if blocking_count
           else f"可复核模板结果 {m['final_profit']} 元。")) for m in report["months"]]
    report["chat_details"].extend(dict(label="报表口径", value=note) for note in report["notes"])
    # Chat is bounded; structured report and XLSX retain every issue and row.
    issue_count = len(report["issues"])
    if issue_count > 500:
        omitted = issue_count - 499
        report["chat_issues"] = report["issues"][:499] + [dict(code="chat_issues_truncated",
            message=f"共 {issue_count} 条核对提示，聊天另有 {omitted} 条未展开；导出报告保留全部问题。",
            blocking=False, severity="info", source_id=None, fact_key=None, month=None)]
    else:
        report["chat_issues"] = list(report["issues"])
    scope = ("、".join(requested) if len(requested) <= 4 else f"{requested[0]} 至 {requested[-1]}（所选 {len(requested)} 个月）") or "当前所选月份"
    report["message"] = ((f"{scope} {cash_pending}，收款、付款和净现金尚不能完整核定。" if missing_cash_months or not report['months'] else
        f"{scope} 已提供流水中的银行收款合计 {report['cumulative']['cash_in']} 元、付款 {report['cumulative']['cash_out']} 元，净现金 {report['cumulative']['net_cash']} 元。")
        + (f"目前有 {blocking_count} 项待确认，完整利润仍为暂算状态，不能用未确认项为零得出利润。" if blocking_count
           else f"已确认业务及业主分成后的模板结果为 {report['cumulative']['final_profit']} 元，可继续复核。"))
    if issue_count > 500:
        report["message"] += f"聊天显示前 499 条核对提示；完整 {issue_count} 条问题已保留在导出报告。"
    return report


# Business labels belong to the presentation layer; calculations keep stable keys.
_CATEGORY_LABELS = {
    "operating_income": "经营收入", "revenue": "经营收入", "ota": "平台结算收入",
    "payroll": "工资", "social_insurance": "五险一金", "bank_fee": "银行手续费",
    "tax": "税费", "laundry_supplier": "布草洗涤费", "cleaning_supplier": "保洁服务费",
    "laundry_supplier_cost": "布草洗涤费", "cleaning_supplier_cost": "保洁服务费",
    "property": "物业费", "maintenance": "维修费", "utility": "水电费",
    "water": "水费", "electricity": "电费", "supplies": "零星采购费", "rent": "租金",
    "operating_expense": "其他经营费用", "owner_distribution": "业主分成",
    "capital": "股东投入", "capital_return": "退还投入", "loan": "借入款",
    "loan_repayment": "归还借款", "internal_transfer": "内部转账",
    "cleaning": "保洁工作记录", "unknown": "待分类",
}
_KIND_LABELS = {"bank": "银行流水", "ota": "平台账单", "cleaning": "保洁记录",
    "utility": "水电物业账单", "image": "截图资料", "bank_transaction": "银行流水",
    "platform_order": "平台结算", "platform_payout": "平台打款",
    "source_expense": "供应方费用", "cleaning_work": "保洁工作", "image_text": "图片文字", "image_expense_candidate": "图片费用候选"}
_CHANNEL_LABELS = {"ctrip": "携程", "douyin": "抖音", "meituan": "美团", "airbnb": "爱彼迎"}
_OWNER_BASIS_LABELS = {"net": "净额（已扣代付）", "gross": "毛额（扣费前）", "unknown": "待确认，暂按净额", "mixed": "净额毛额混合，待核实", "missing": "缺少当月分成依据", "not_applicable": "未列分成与代付"}
_ISSUE_GUIDANCE = {
    "counterpart_business_month_conflict": ("关联事项月份冲突", "核实各份来源是否为同一事项，统一业务归属月份，保留银行实际到账日期。"),
    "counterpart_category_conflict": ("关联事项类别冲突", "先核实是否同一笔业务，再统一类别；本金、费用以及明确水费和电费不能混作同一项。"),
    "owner_distribution_basis_unknown": ("业主分成口径待确认", "确认当月分成是扣除代付后的净额，还是扣费前毛额；未知时暂按净额计算，不能定稿。"),
    "owner_distribution_basis_mixed": ("业主分成混合口径", "统一当月分成口径，或补齐逐业主代付分配依据，避免重复扣减或漏扣代付。"),
    "owner_settlement_missing": ("代付缺少当月结算依据", "补充当月业主分成金额，并明确净额已扣代付或毛额包含代付；空白不代表无需分成。"),
    "image_amount_unconfirmed": ("图片费用待核实", "逐项核实图片金额、用途、期间、承担方及付款方；不采用的候选请明确排除。"),
    "image_payment_unconfirmed": ("图片实际付款方待确认", "确认这笔费用实际由公司还是业主付款。"),
    "image_date_unconfirmed": ("图片费用日期待确认", "确认费用日期或明确按业务月末归集；不得把截图时间当付款日期。"),
    "image_discount_unconfirmed": ("图片优惠待分摊", "确认按各房原金额比例分摊优惠，或采用实付汇总并排除房间明细。"),
    "image_owner_allocation_required": ("业主费用待分房", "核实各房费用及优惠分摊后，再归入各业主。"),
    "image_alternative_conflict": ("图片汇总与明细重复", "仅保留汇总或明细一种口径，明确排除另一种。"),
    "classification_unconfirmed": ("业务分类待确认", "确认收支类别及业务归属月份；费用还需确认承担方和实际付款方。"),
    "category_unresolved": ("业务分类待确认", "将这笔款项归入明确的收入、费用、分成或非经营类别。"),
    "business_month_unknown": ("归属月份待确认", "依据订单或费用所属期间确认月份，不直接使用银行到账月份。"),
    "expense_payer_unconfirmed": ("费用承担方待确认", "确认公司承担还是业主承担，并说明实际由谁付款。"),
    "owner_expense_payment_unknown": ("实际付款方待确认", "确认业主是否已付，或公司是否垫付，并核对业主结算扣除。"),
    "possible_duplicate_business_evidence": ("同笔款项待对应", "将图片、费用表与银行流水对应，确认是否为同笔款项，避免重复计入。"),
    "counterpart_amount_mismatch": ("对应金额不一致", "核对退款、手续费或分笔收付，补足金额差异的依据。"),
    "counterpart_payment_conflict": ("付款方与流水冲突", "核对是否公司垫付、报销或对应错误，再确认实际付款方。"),
    "counterpart_payer_conflict": ("关联费用承担方不一致", "逐项核对关联来源，统一确认公司承担或业主承担，不用付款人代替承担方。"),
    "mixed_counterpart_direction": ("收付款对应待拆分", "将收款与付款分别对应，说明各笔资金用途。"),
    "category_direction_conflict": ("收付方向与分类冲突", "核对是否退款或冲销，再确认收支分类。"),
    "missing_linked_evidence": ("缺少对应资料", "补充已关联款项的原始资料，再重新核对。"),
    "invalid_evidence_links": ("资料对应关系待修正", "重新选择这笔款项对应的原始资料。"),
    "no_evidence": ("缺少原始资料", "上传需要核对的原始账单和银行流水。"),
    "bank_evidence_missing": ("缺少银行流水", "补充所选月份的完整银行流水，用于核对公司整体收付款。"),
    "requested_month_evidence_missing": ("所选月份缺少资料", "补充该月流水及业务明细；缺少资料不代表收入和费用为零。"),
    "business_amount_unknown": ("金额待核实", "对照原表核实金额，无法确认时补充原始凭据。"),
    "invalid_cash_fact": ("银行金额或方向异常", "对照原流水核实金额与收付方向后重新导入。"),
    "cash_month_unknown": ("银行日期待核实", "补充完整交易日期，才能归入月度现金统计。"),
    "image_requires_verification": ("图片文字待核实", "核实图片中的金额、所属月份、承担方以及与表格的对应关系。"),
    "ocr_unavailable": ("图片内容待补充", "补充清晰原图、文字或原始表格，再核实金额与月份。"),
    "ocr_failed": ("图片内容待补充", "补充清晰原图、文字或原始表格，再核实金额与月份。"),
    "ocr_low_coverage": ("图片内容待补充", "图片文字可能不完整，请补充原始表格或逐项文字。"),
    "ctrip_cash_unconfirmed": ("平台结算与银行实收待对应", "按离店日核对业务月份，再对应银行实际到账，不重复计算。"),
    "douyin_cash_basis": ("平台结算与银行实收待对应", "核对平台明细、打款与银行到账的对应关系及业务所属月份。"),
    "cleaning_evidence_only": ("保洁工作与费用分别核对", "根据实际工作记录和约定收费核对，不能直接按表列次数推算费用。"),
    "owner_borne_expense_excluded": ("业主承担费用已单列", "核对业主结算中的费用扣除；本笔不作为公司自身成本。"),
    "confirmed_duplicate_excluded": ("重复款项已排除", "已按确认的对应关系只计入一次，保留原始凭据。"),
    "duplicate_source_fact": ("重复资料已去重", "同一来源明细只计入一次，无需重复记账。"),
    "worksheet_dimensions_recovered": ("原表读取说明", "已恢复原表实际行列；保留原文件供复核。"),
}


def export_workbook(report: dict, sources: list[dict]) -> bytes:
    """Human-readable, literal XLSX presentation; never changes report arithmetic."""
    import math
    from datetime import date, datetime
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    font_name = "Microsoft YaHei"
    ink, stone, linen, sand = "2B2721", "6B665B", "E5DDCB", "F5F1EA"
    money_format = '#,##0.00;(#,##0.00);"—"'
    source_map = {s.get("source_id"): s for s in sources}
    for item in report.get("source_inventory", []):
        source_map.setdefault(item.get("source_id"), item)
    originals = {f.get("key"): f for s in sources for f in (s.get("parsed") or {}).get("facts", [])}
    details = {r.get("fact_key"): r for r in report.get("details", [])}
    blocked_keys = {i.get("fact_key") for i in report.get("issues", []) if i.get("blocking")}
    selected = {r.get("fact_key"): r for r in report.get("profit_rows", [])}
    months = [m["month"] for m in report.get("months", [])]
    scope = "、".join(months) or "未指定月份"

    def clean(value, default="—"):
        if value is None or value == "":
            return default
        text = str(value)
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        text = re.sub(r"\b(?:FCS|FCP|FCR|MCR|EXP|DOC|TSK|USR)-[A-Za-z0-9-]+\b|\b[0-9a-fA-F]{64}\b", "相关资料", text)
        if text.lstrip().startswith(("=", "+", "-", "@")):
            text = "'" + text
        return text[:32767]

    def chinese(value, default="待核实"):
        text = clean(value, default)
        for key, label in sorted({**_CATEGORY_LABELS, **_KIND_LABELS, **_CHANNEL_LABELS,
            "checkout_date": "离店日期", "platform_settlement_date": "平台结算日期",
            "platform_payout_date": "平台打款日期", "provider_collection_date": "供应方收款日期",
            "service_date": "服务日期", "company_advanced": "公司垫付", "company": "公司",
            "owner_paid": "业主已付", "owner": "业主", "OTA": "平台", "OCR": "图片文字识别"}.items(), key=lambda item: -len(item[0])):
            text = re.sub(r"(?<![A-Za-z_])" + re.escape(key) + r"(?![A-Za-z_])", label, text)
        # Technical fallback codes and JSON do not become reader-facing prose.
        if not re.search(r"[\u3400-\u9fff]", text) or text.lstrip().startswith(("{", "[")):
            return default
        return text

    def amount(value):
        parsed = _dec(value) if value is not None else None
        return parsed if parsed is not None else "待确认"

    def day(value):
        if isinstance(value, (date, datetime)):
            return value
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return "待确认"

    def worksheet_name(value):
        text = clean(value)
        default = re.fullmatch(r"Sheet\s*(\d+)", text, re.IGNORECASE)
        if default:
            return f"工作表{default.group(1)}"
        if text.lower() in {"july", "jul"}:
            return "七月表"
        return chinese(text, "原工作表") if re.search(r"[A-Za-z]", text) else text

    def position(ref):
        source = source_map.get(ref.get("source_id"), {})
        filename = clean(source.get("filename"), "来源文件待补充")
        sheet_name = worksheet_name(ref.get("sheet"))
        row = ref.get("row1based", ref.get("row"))
        return filename, sheet_name, row

    def provenance(row):
        references = row.get("lineage") or []
        if not references:
            original = originals.get(row.get("fact_key"), {})
            references = [dict(original, source_id=row.get("source_id"))]
        grouped = {}
        for reference in references:
            filename, sheet_name, index = position(reference)
            grouped.setdefault((filename, sheet_name), set())
            if isinstance(index, int):
                grouped[(filename, sheet_name)].add(index)
        result = []
        for (filename, sheet_name), indexes in grouped.items():
            spans = []
            for index in sorted(indexes):
                if spans and index == spans[-1][-1] + 1:
                    spans[-1].append(index)
                else:
                    spans.append([index])
            row_text = "、".join(str(span[0]) if len(span) == 1 else f"{span[0]}至{span[-1]}" for span in spans)
            result.append(f"{filename}\n{sheet_name}" + (f"，第{row_text}行" if row_text else "，原位置待补充"))
        return "\n".join(result)

    def style(ws, widths, header_row=5, numeric_columns=()):
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f"B{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.print_title_rows = f"1:{header_row}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.25
        ws.page_margins.top = ws.page_margins.bottom = 0.4
        ws.oddHeader.center.text = ws.title
        ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
        ws.print_area = ws.dimensions
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in ws:
            for cell in row:
                cell.font = Font(name=font_name, size=11, color=ink)
                cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
                if isinstance(cell.value, (Decimal, int, float)):
                    cell.alignment = Alignment(vertical="center", horizontal="right")
                    if cell.column in numeric_columns:
                        cell.number_format = money_format
                elif isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
                    cell.alignment = Alignment(vertical="center", horizontal="center")
            if row[0].row > header_row:
                heights = [max(1, sum(max(1, math.ceil(sum(2 if ord(ch) > 255 else 1 for ch in line) / max(8, widths[c.column - 1] - 2))) for line in str(c.value or "").split("\n"))) for c in row]
                ws.row_dimensions[row[0].row].height = max(28, max(heights) * 17)
        ws.row_dimensions[1].height = 8
        ws.row_dimensions[2].height = 30
        ws.row_dimensions[3].height = 24
        ws.row_dimensions[4].height = 26
        ws.row_dimensions[header_row].height = 30
        ws.cell(2, 1).font = Font(name=font_name, size=16, color=ink)
        for cell in ws[header_row]:
            cell.fill = PatternFill("solid", fgColor=ink)
            cell.font = Font(name=font_name, size=11, color="FBF8F1")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_num in (2, 3, 4):
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=ws.max_column)
        ws.cell(3, 1).font = Font(name=font_name, size=11, color=stone)
        ws.cell(4, 1).font = Font(name=font_name, size=11, color="8A6E5A")

    def table(name, headers, rows, widths, note, numeric_columns=()):
        ws = workbook.create_sheet(name)
        ws.append([])
        ws.append([name])
        ws.append([f"报告范围：{scope}　金额单位：元"])
        ws.append([note])
        ws.append(headers)
        for row in rows:
            ws.append([value if isinstance(value, (Decimal, int, float, date, datetime)) and not isinstance(value, bool) else clean(value) for value in row])
        style(ws, widths, numeric_columns=numeric_columns)
        return ws

    # The template's business sequence comes first. Cash is a separate section.
    summary = workbook.create_sheet("经营概览")
    summary.append([])
    summary.append(["经营收支与利润"])
    summary.append([f"报告范围：{scope}　金额单位：元"])
    blocking = report.get("metrics", {}).get("open_issue_count", 0)
    summary.append([f"{blocking} 项待处理，完整利润暂不能定稿。下列业务金额仅含已确认部分。" if blocking else "业务结果已具备复核条件；请核对来源和实际账务后使用。"])
    columns = report.get("months", []) + [report.get("cumulative", {})]
    summary.append(["项目"] + months + ["所选月份累计"])
    sections, totals = [], []

    def band(label):
        summary.append([label]); sections.append(summary.max_row)

    def confirmed_value(column, key, period=None):
        role = {"confirmed_revenue": "revenue", "confirmed_expenses": "expense", "confirmed_owner_share": "owner_share", "confirmed_owner_advances": "owner_advance"}.get(key)
        has_confirmed = any(r.get("business_month") in (months if period is None else [period]) and (role is None or r.get("role") == role) for r in report.get("profit_rows", []))
        if key == "final_profit" or not blocking or has_confirmed:
            return amount(column.get(key))
        return "待确认"

    def result_row(label, key, total=False, business=True):
        values = [confirmed_value(column, key, months[i] if i < len(months) else None) if business else
                  amount(column.get(key)) if column.get(key) is not None else "缺少银行流水" for i, column in enumerate(columns)]
        summary.append([label] + values)
        if total:
            totals.append(summary.max_row)

    band("一、经营收入（按业务月份）")
    result_row("已确认经营收入", "confirmed_revenue", True)
    band("二、公司经营费用（按业务月份）")
    supplier_category = lambda value: {"cleaning_supplier_cost": "cleaning_supplier", "laundry_supplier_cost": "laundry_supplier"}.get(value, value)
    present_categories = {supplier_category(r.get("category")) for r in report.get("details", [])} & EXPENSES
    for category in ["payroll", "social_insurance", "bank_fee", "tax", "supplies", "cleaning_supplier", "laundry_supplier", "water", "electricity", "utility", "property", "maintenance", "rent", "operating_expense"]:
        if category not in present_categories:
            continue
        values = []
        for period in months + [None]:
            rows = [r for r in report.get("profit_rows", []) if r.get("role") == "expense" and supplier_category(r.get("category")) == category and r.get("business_month") in (months if period is None else [period])]
            pending = any(supplier_category(r.get("category")) == category and r.get("include", True) and (not r.get("confirmed") or r.get("fact_key") in blocked_keys) and (r.get("business_month") is None or r.get("business_month") in (months if period is None else [period])) for r in report.get("details", []))
            has_month_evidence = any(r.get("business_month") in (months if period is None else [period]) for r in report.get("profit_rows", []))
            values.append(sum((_dec(r.get("amount")) or ZERO for r in rows), ZERO) if rows or not (pending or (blocking and not has_month_evidence)) else "待确认")
        summary.append([_CATEGORY_LABELS[category]] + values)
    result_row("公司经营费用小计", "confirmed_expenses", True)
    band("三、业主分成（单独列示）")
    result_row("已确认业主分成", "confirmed_owner_share", True)
    result_row("代业主支付费用", "confirmed_owner_advances")
    summary.append(["业主分成计算口径"] + [_OWNER_BASIS_LABELS.get(m.get("owner_distribution_basis"), "待确认") for m in report.get("months", [])] + ["按各月口径分别计算"])
    band("四、经营结果")
    result_row("分成前经营结余", "confirmed_operating_result")
    result_row("已确认部分暂算利润", "provisional_template_result")
    result_row("可复核利润", "final_profit", True)
    band("五、银行现金（按实际交易日期）")
    result_row("银行收款", "cash_in", business=False)
    result_row("银行付款", "cash_out", business=False)
    result_row("银行净现金", "net_cash", True, business=False)
    for direction, label in (("in", "其中：用途待确认收款"), ("out", "其中：用途待确认付款")):
        summary.append([label] + [sum((_dec(r.get("amount")) or ZERO for r in report.get("cash_rows", []) if not r.get("confirmed") and r.get("direction") == direction and r.get("cash_month") in (months if period is None else [period])), ZERO)
            if columns[i].get("cash_in") is not None else "缺少银行流水" for i, period in enumerate(months + [None])])
    summary.append(["银行净现金含借款、股东投入等资金往来，不等同于利润。待确认项不能按零处理。"])
    style(summary, [32] + [20] * len(columns), numeric_columns=tuple(range(2, len(columns) + 2)))
    summary.auto_filter.ref = None
    summary.sheet_properties.tabColor = ink
    summary.row_dimensions[summary.max_row].height = 30
    summary.merge_cells(start_row=summary.max_row, start_column=1, end_row=summary.max_row, end_column=summary.max_column)
    for row_number in sections:
        summary.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=summary.max_column)
        for cell in summary[row_number]:
            cell.fill = PatternFill("solid", fgColor=sand)
        summary.row_dimensions[row_number].height = 26
    for row_number in totals:
        for cell in summary[row_number]:
            cell.border = Border(top=Side(style="thin", color=linen), bottom=Side(style="double", color=linen))
    # Large requested windows print as successive readable month blocks.
    if len(months) > 6:
        summary.page_setup.fitToWidth = 0
        summary.page_setup.scale = 85
        summary.print_title_cols = "A:A"
        for index in range(7, summary.max_column, 6):
            summary.col_breaks.append(Break(id=index))

    def row_values(row):
        original = originals.get(row.get("fact_key"), {})
        kind = row.get("kind")
        role = selected.get(row.get("fact_key"), {}).get("role")
        if not row.get("include", True):
            status = "已确认排除，不重复计入"
        elif role:
            status = {"revenue": "计入已确认收入", "expense": "计入已确认公司费用", "owner_share": "计入已确认业主分成", "owner_advance": "代业主支付费用，单独列示"}.get(role, "已确认")
        elif row.get("confirmed"):
            status = "已确认，未计入公司利润" if row.get("category") in NON_OPERATING or row.get("payer") in {"owner", "业主"} else "已确认，仍需核对对应关系"
        else:
            status = "图片文字待核实" if kind == "image_text" else "工作证据，费用另核" if kind == "cleaning_work" else "分类及业务归属待确认"
        nature = "银行收款" if kind == "bank_transaction" and row.get("direction") == "in" else "银行付款" if kind == "bank_transaction" and row.get("direction") == "out" else "平台结算净额（非银行实收）" if kind in {"platform_order", "platform_payout"} else "供应方账单金额" if kind == "source_expense" else "工作记录，不推导费用" if kind == "cleaning_work" else "图片文字，不作为入账金额"
        if kind == "image_expense_candidate":
            nature = "图片费用金额（非银行实付）"
            if row.get("fields", {}).get("applied_discount_allocation") == "proportional":
                nature = "按已确认优惠比例分摊金额"
            if row.get("fact_key") in blocked_keys:
                status = "图片费用仍待核实，未定稿"
        elif kind == "image_text" and row.get("candidate_evidence_only"):
            status = "图片原文，见逐项费用候选"
        category = _CATEGORY_LABELS.get(row.get("category"), "待分类")
        channel = _CHANNEL_LABELS.get(row.get("channel"))
        if channel:
            category = f"{channel}结算收入"
        room = clean(row.get("room_ref", original.get("room_ref")))
        if kind == "bank_transaction":
            description = category  # Never re-export private bank narratives/IDs.
        elif kind == "image_text":
            description = "请核实原图中的业务信息"
        elif kind == "image_expense_candidate":
            description = "图片费用依据"
        else:
            description = chinese(row.get("description"), category)
        if kind in {"source_expense", "image_expense_candidate"} or kind == "bank_transaction" and row.get("payer"):
            description += f"\n承担方：{chinese(row.get('payer'), '待确认')}\n实际付款方：{chinese(row.get('paid_by'), '待确认')}"
            if row.get("independent_cost"):
                description += "\n已明确核实为独立费用"
        if row.get("category") == "owner_distribution":
            description += "\n分成口径：" + _OWNER_BASIS_LABELS.get(row.get("owner_distribution_basis"), "待确认，暂按净额")
        if kind == "image_expense_candidate" and row.get("date_basis") == "month_end_accrual":
            description += "\n按业务月末归集，原付款日期：" + clean(row.get("payment_date"), "待核实")
        if kind == "image_expense_candidate" and row.get("fields", {}).get("original_amount") is not None:
            description += "\n优惠前金额：" + clean(row["fields"]["original_amount"]) + " 元"
        business_month = clean(row.get("business_month"), "待确认") + ("（待确认）" if row.get("business_month") and not row.get("confirmed") else "")
        return [day(row.get("date")), business_month, category, "—" if kind in {"cleaning_work", "image_text"} else amount(row.get("amount")), nature, status, f"{room}\n{description}" if room != "—" else description, provenance(selected.get(row.get("fact_key"), row))]

    table("收支与工作明细", ["日期", "业务月份", "收支类别", "金额（元）", "金额性质", "核对状态", "房间或事项", "来源文件与原表位置"],
        (row_values(row) for row in report.get("details", [])), [15, 18, 18, 17, 26, 28, 28, 52],
        "保留原始来源的全部明细；经营概览仅汇总所选月份。平台账单与银行款项不能直接相加。", (4,))

    def issue_info(item):
        code = str(item.get("code") or "")
        title, action = _ISSUE_GUIDANCE.get(code, ("原始资料待核实", "对照原文件核实提示所指内容，必要时补充资料或调整读取方式。"))
        row = details.get(item.get("fact_key"), {})
        ref = row or dict(source_id=item.get("source_id"), lineage=[])
        if not row:
            original_issue = next((issue for issue in (source_map.get(item.get("source_id"), {}).get("parsed") or {}).get("issues", []) if issue.get("code") == code), {})
            ref["lineage"] = [dict(original_issue, source_id=item.get("source_id"))]
        message = chinese(item.get("message"), title)
        return title, action, row, provenance(ref), message

    grouped = {}
    issue_rows = []
    for item in report.get("issues", []):
        title, action, row, reference, message = issue_info(item)
        key = (title, action, bool(item.get("blocking")))
        group = grouped.setdefault(key, dict(count=0, sources=set()))
        group["count"] += 1
        group["sources"].add(clean(source_map.get(item.get("source_id"), {}).get("filename"), "整体报告"))
        issue_rows.append([title, "影响完整利润" if item.get("blocking") else "核对说明", clean(item.get("month"), "待确认"), message, action, reference])
    table("待处理事项", ["事项", "数量", "影响", "下一步怎么做", "涉及文件"],
        ([key[0], group["count"], "处理后才能复核利润" if key[2] else "核对说明", key[1], "\n".join(sorted(group["sources"]))] for key, group in sorted(grouped.items(), key=lambda entry: (not entry[0][2], entry[0][0]))),
        [28, 10, 28, 62, 62], "按事项合并展示；每条具体问题和原表位置见“问题明细”。")
    table("问题明细", ["事项", "影响", "业务月份", "具体情况", "下一步怎么做", "来源文件与原表位置"],
        issue_rows, [28, 19, 15, 64, 60, 52], "逐项保留所有核对提示；可按事项、影响和月份筛选。")
    table("资料清单", ["原始文件", "资料类型", "明细数量", "读取提示数量", "使用口径"],
        ([clean(item.get("filename")), _KIND_LABELS.get(item.get("kind"), "其他原始资料"), item.get("fact_count", 0), item.get("issue_count", 0), "图片文字待核实，不代表可入账明细" if item.get("kind") == "image" else "保洁工作记录，费用另行核对" if item.get("kind") == "cleaning" else "供应方收款不证明由公司承担或已垫付" if item.get("kind") == "utility" else "银行实际收付，业务分类和归属需确认" if item.get("kind") == "bank" else "平台结算须对应银行实收"] for item in report.get("source_inventory", [])),
        [64, 24, 16, 18, 60], "原始文件名和原表位置用于查找依据；图片文字条数不能作为财务记录数量。")
    table("原利润表对照", ["原表项目", "原表缓存金额（元）", "来源文件与原表位置", "当前口径"],
        ([chinese(entry.get("label"), "原表项目待核实"), amount(entry.get("cached_amount")), provenance(dict(source_id=template.get("source_id"), lineage=[dict(source_id=template.get("source_id"), sheet=template.get("sheet"), row1based=entry.get("row"))])), "原利润表对照值，未并入本次核对结果"] for template in report.get("templates", []) for entry in template.get("entries", [])),
        [44, 26, 62, 62], "保留原七月利润表的收入、支出、分成和利润项目，作为对照；不是本次已确认结果。", (2,))
    table("口径说明", ["口径", "说明"], [
        ["经营收入", "按已确认的业务月份统计；平台净结算与银行到账可能对应同一笔收入，先对应再计入。"],
        ["公司经营费用", "仅列入已确认由公司承担的费用；业主承担、公司垫付和实际付款方分别核实。"],
        ["业主分成", "分成净额指已扣除代付后的净付或净应付；暂算利润还要另减代业主支付费用。分成毛额指扣费前金额，已包含代付，不再重复扣减。"],
        ["代业主支付费用", "仅纳入已确认业主承担且公司实际支付的费用，单独列示，不混入公司自身费用。未知、混合分成口径或缺当月结算依据时，完整利润仍待核实。"],
        ["银行现金", "按交易日统计收款、付款和净额；借款、股东投入与内部转账不能直接当作收入或费用。"],
        ["待确认与零", "待确认表示资料或处理依据不完整，不能按零处理。数字零仅表示当前已确认部分的合计为零。"],
        ["暂算与可复核", "暂算只涵盖已确认部分；影响完整利润的问题未处理完时，可复核利润保持待确认。"],
        ["系统费用台账", "系统服务收费金额不是供应方实际成本；不在本报告重复计入。需另查系统台账并核对对应关系。"],
        ["来源位置", "明细列出文件名、工作表及原行号。默认工作表名按编号显示；连续原行号合并成范围，对应资料仍逐份保留。"],
        ["原利润表", "原利润表缓存金额仅作对照；不会自动改变本次经营结果。"],
    ], [26, 120], "金额使用数值单元格，日期可排序；所有来源文本均为普通文本，不执行原表公式。")

    # A reader should get an answer before encountering the accounting schedules.
    # This page explains existing values; it never promotes parser suggestions
    # into confirmed expenses or treats a cash difference as profit/bank balance.
    overview = workbook.create_sheet("先看这页", 0)
    latest = report.get("months", [])[-1] if report.get("months") else {}
    latest_month = latest.get("month", "")
    month_label = (f"{latest_month[:4]}年{int(latest_month[5:])}月" if latest_month else "本次")
    has_cash = bool(latest.get("cash_fact_keys"))
    overview.append([])
    overview.append([f"{month_label}，这笔账现在是什么情况"])
    overview.append([f"本报告覆盖 {len(months)} 个月；此页先看最近的{month_label}，各月对照在后面。"])
    overview.append(["金额单位：元。以下是你提供的资料所反映的情况。"])
    overview.append(["你关心的事", "现在的结果", "怎么理解"])
    overview.append(["这个月到底赚了多少", amount(latest.get("final_profit")) if latest.get("final_profit") is not None else "暂时还算不准",
                     "收入、费用和业主分成都已核对，可查看后面的计算明细。" if latest.get("final_profit") is not None else
                     "还要分清房费、投资款、业主分成及重复单据。上传成功不表示这些账已经核实。"])
    overview.append(["银行一共收进来", amount(latest.get("cash_in")) if has_cash else "缺少该月流水", "银行收到的钱；可能包含投资款、借款，不能全当营业收入。"])
    overview.append(["银行一共付出去", amount(latest.get("cash_out")) if has_cash else "缺少该月流水", "银行付出的钱；可能包含业主分成、垫付款及其他往来。"])
    overview.append(["收进来减去付出去", amount(latest.get("net_cash")) if has_cash else "暂时无法计算", "只表示这份流水的收付差额，不是利润，也不是银行卡余额。"])
    overview.append([])
    overview.append(["付款大致花在哪里"])
    overview.append(["下面按流水描述初步分组，用于找账；分类还需核实，不能直接当成公司成本。"])
    overview.append(["付款去向（初步分类）", "银行付款金额", "需要注意什么"])
    spending = sorted(((category, _dec(values.get("out")) or ZERO)
                       for category, values in latest.get("cash_categories", {}).items()
                       if (_dec(values.get("out")) or ZERO) != ZERO), key=lambda item: -item[1])
    # Keep the first page short even when a bank contains many cost categories.
    shown = spending[:6] if len(spending) > 7 else spending
    for category, value in shown:
        overview.append(["还没分清用途" if category == "unknown" else _CATEGORY_LABELS.get(category, "其他付款"), value,
                         "先补充这些付款的用途，再判断是否属于经营费用。" if category == "unknown" else
                         "分类待核实"])
    if len(spending) > 7:
        overview.append(["其他付款合计", sum((value for _, value in spending[6:]), ZERO), "具体项目可在后面的收支明细中查找。"])
    if not spending:
        overview.append(["付款去向", "暂无可列明金额", "没有付款明细可供分类，不据此判断费用为零。"])
    overview.append([])
    overview.append(["接下来，先确认这些事"])
    action_heading = overview.max_row
    codes = {str(item.get("code")) for item in report.get("issues", []) if item.get("blocking")}
    actions = []
    if "classification_unconfirmed" in codes or "business_month_unknown" in codes:
        unknown = latest.get("cash_categories", {}).get("unknown", {})
        unknown_out, unknown_in = _dec(unknown.get("out")) or ZERO, _dec(unknown.get("in")) or ZERO
        amounts = "；".join(f"{label} {value:,.2f} 元" for label, value in [("用途不明的付款", unknown_out), ("用途不明的收款", unknown_in)] if value)
        actions.append(["先分清款项的用途", amounts or "已有初步分类，尚待核实", "说明哪些是房费、股东投入、借款、公司费用或业主分成，并确认属于哪个月。"])
    if any("duplicate" in code or "counterpart" in code or "cash_unconfirmed" in code for code in codes):
        actions.append(["同一笔钱只算一次", "对照银行、平台和费用单", "核对平台打款是否已在银行流水里；费用表和付款记录若是同一笔，要关联起来。"])
    if any("image" in code for code in codes):
        actions.append(["核实截图里的费用", "确认金额、月份和付款人", "查看物业、维修等原图；遇到优惠、合计与分房明细，先确认用哪一组金额。"])
    if any("payer" in code or "payment" in code or "owner" in code for code in codes):
        actions.append(["确认谁承担费用", "公司费用还是业主费用", "说明实际由谁支付；业主分成是扣除垫付款后的金额，还是扣除前的金额。"])
    if blocking and not actions:
        actions.append(["核对剩余差异", "见后面的待处理事项", "按每类问题核对原始资料；这里没有将缺少的依据当作零。"])
    for action in actions:
        overview.append(action)
    if not blocking:
        overview.append(["本次核对", "已具备复核条件", "核对后面的收支明细，再用于结账。"])
    overview.append([])
    overview.append(["查看具体待办 →", "查看每笔明细 →", "查看各月对照 →"])
    navigation_row = overview.max_row
    overview.append(["后面的表供核对细节使用；银行到账、系统收费和公司利润分别看，不能直接相加。"])
    style(overview, [29, 27, 68], numeric_columns=(2,))
    overview.auto_filter.ref = None
    overview.freeze_panes = "A6"
    overview.page_setup.paperSize = overview.PAPERSIZE_A4
    overview.page_setup.orientation = "portrait"
    overview.page_setup.fitToHeight = 1
    overview.sheet_properties.tabColor = "7B8578"
    for number in (11, 12, action_heading, overview.max_row):
        overview.merge_cells(start_row=number, start_column=1, end_row=number, end_column=3)
        overview.row_dimensions[number].height = 30
    for number in (11, action_heading):
        overview.cell(number, 1).fill = PatternFill("solid", fgColor=sand)
        overview.cell(number, 1).font = Font(name=font_name, size=13, color=ink)
    for cell in overview[13]:
        cell.fill = PatternFill("solid", fgColor=sand)
    overview.row_dimensions[6].height = 62
    overview.cell(6, 2).font = Font(name=font_name, size=16, color=ink)
    for number in (7, 8, 9):
        overview.row_dimensions[number].height = 48
        overview.cell(number, 2).font = Font(name=font_name, size=15, color=ink)
    for column, target in enumerate(("待处理事项", "收支与工作明细", "经营概览"), 1):
        cell = overview.cell(navigation_row, column)
        cell.hyperlink = f"#'{target}'!A1"
        cell.font = Font(name=font_name, size=11, color=stone, underline="single")
    workbook.active = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
