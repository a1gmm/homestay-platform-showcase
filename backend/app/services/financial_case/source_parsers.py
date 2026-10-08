"""Deterministic, read-only source evidence; no ledger or payment assertions.

Facts are the one countable detail level. Summaries are independent controls or
alternate views, never additional revenue/expense facts. No guest/customer,
account, phone, voucher or staff identity cells are retained.
"""
from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256

from app.services.billing_recon.parser import BillParseError, load_workbook_rows


def _text(value):
    return str(value).strip() if value is not None else ""


def _decimal(value):
    try:
        result = Decimal(_text(value).replace(",", "").replace("¥", ""))
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def _money(value):
    return format(value.quantize(Decimal("0.01")), "f") if value is not None else None


def _day(value, datemode=0):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and value > 0:
        import xlrd
        try:
            return xlrd.xldate_as_datetime(value, datemode).date().isoformat()
        except (ValueError, OverflowError, xlrd.XLDateError):
            return None
    match = re.match(r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})(?:日|\b)", _text(value))
    if match:
        try:
            return date(*map(int, match.groups())).isoformat()
        except ValueError:
            pass
    return None


def _header(rows, required):
    for index, row in enumerate(rows[:30]):
        cells = ["".join(_text(v).split()) for v in row]
        if required <= set(cells):
            return index, {name: i for i, name in enumerate(cells) if name}
    return None


def _cell(row, headers, name):
    col = headers.get(name)
    return row[col] if col is not None and col < len(row) else None


def _issue(result, code, message, sheet=None, row=None, severity="warning"):
    result["issues"].append(dict(code=code, message=message, sheet=sheet,
                                 row1based=row, severity=severity))


def _fact(digest, sheet, row, kind, amount=None, day=None, suffix="", **fields):
    key = sha256(f"{digest}:{sheet}:{row}:{kind}:{suffix}".encode()).hexdigest()
    return dict(key=key, sheet=sheet, row1based=row, kind=kind, date=day,
                business_month=day[:7] if day else None, amount=_money(amount),
                description="", room_ref=None, **fields)


def _load(data, filename):
    sheets, datemode = load_workbook_rows(data, filename)
    repaired = []
    # Some exports falsely declare A1:A1. Preserve the shared loader, but recover
    # these sheets by streaming actual XML rows, with the same row limit. Douyin
    # needs column 65 (settlement time), beyond the generic loader's 64 columns.
    if not filename.lower().endswith(".xls") and (
        any(len(rows) == 1 and len(rows[0]) == 1 for rows in sheets.values())
        or "货款结算-正向-团购" in sheets
    ):
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                old = sheets[sheet.title]
                recover = len(old) == 1 and len(old[0]) == 1
                if not recover and sheet.title != "货款结算-正向-团购":
                    continue
                sheet.reset_dimensions()
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) >= 20000:
                        raise BillParseError("表格行数超限(>20000)")
                    rows.append(list(row[:128]))
                if recover and rows != old:
                    repaired.append(sheet.title)
                sheets[sheet.title] = rows
        finally:
            workbook.close()
    return sheets, datemode, repaired


def _identifier(value, result, sheet, row):
    if isinstance(value, (float, int)):
        if abs(value) >= 10**15 or (isinstance(value, float) and not value.is_integer()):
            _issue(result, "numeric_order_id_precision", "数字型长订单号可能已丢失精度，不能据此自动匹配", sheet, row, "error")
            return None
        return str(int(value))
    text = _text(value)
    if re.fullmatch(r"\d+(?:\.0)?[eE][+-]?\d+", text):
        _issue(result, "scientific_order_id", "科学计数法订单号需原始文本核实", sheet, row, "error")
        return None
    return text or None


def _utility(result, sheets, digest, datemode):
    handled = set()
    totals = defaultdict(Decimal)
    for sheet, rows in sheets.items():
        found = _header(rows, {"房间名称", "收款日期", "费用科目", "已收金额"})
        if not found:
            continue
        handled.add(sheet)
        header, cols = found
        for index, row in enumerate(rows[header + 1:], header + 2):
            if not any(_text(v) for v in row):
                continue
            label = _text(_cell(row, cols, "费用科目"))
            if any(_text(v) in {"合计", "总计", "小计"} for v in row):
                result["summaries"].append(dict(kind="utility_reported_total", sheet=sheet,
                    row1based=index, amount=_money(_decimal(_cell(row, cols, "已收金额")))))
                continue
            category = next((kind for word, kind in [("水费", "water"), ("电费", "electricity"), ("物业", "property")] if word in label), "unknown")
            amount = _decimal(_cell(row, cols, "已收金额"))
            day = _day(_cell(row, cols, "收款日期"), datemode)
            fact = _fact(digest, sheet, index, "source_expense", amount, day,
                         category=category, fields={"provider_collected_amount": _money(amount),
                         "payer": "unconfirmed", "expense_ownership": "unconfirmed",
                         "date_basis": "provider_collection_date"})
            fact.update(description=label, room_ref=_text(_cell(row, cols, "房间名称")) or None)
            result["facts"].append(fact)
            if amount is None or category == "unknown" or day is None:
                _issue(result, "utility_row_needs_review", "费用类别、金额或收款日期无法完整识别", sheet, index, "error")
            if amount is not None:
                totals[category] += amount
    result["summaries"].append(dict(kind="utility_detail_total", amount=_money(sum(totals.values(), Decimal(0))),
        categories={k: _money(v) for k, v in totals.items()}, count=len(result["facts"]),
        room_count=len({f["room_ref"] for f in result["facts"] if f["room_ref"]})))
    _issue(result, "expense_payer_unconfirmed", "已收金额是物业/供应方收款视角；公司是否垫付、费用归属及是否已记账需确认")
    return handled


def _ctrip(result, sheets, digest, datemode):
    handled = set()
    orders = {}
    raw_count = 0
    nights = Decimal(0)
    controls = []
    money_fields = {"gross": "订单折前底价", "discount": "底价折扣金额", "subsidy": "底价补款金额", "miscellaneous": "闪住杂费"}
    for sheet, rows in sheets.items():
        found = _header(rows, {"订单号", "结算价", "离店日期", "间夜"})
        if not found:
            if sheet == "折扣金额明细" and _header(rows, {"订单号", "折扣金额"}):
                handled.add(sheet)
                result["summaries"].append(dict(kind="discount_detail_reference", sheet=sheet,
                    countable=False, note="折扣分解明细不重复计入净结算额"))
            control_header = _header(rows, {"结算周期", "结算价", "间夜数"})
            if control_header:
                handled.add(sheet)
                h, cols = control_header
                for i, row in enumerate(rows[h + 1:], h + 2):
                    period = _text(_cell(row, cols, "结算周期"))
                    if not re.fullmatch(r"\d{8}-\d{8}", period):
                        continue
                    amount = _decimal(_cell(row, cols, "结算价"))
                    control = dict(kind="ctrip_reported_total", sheet=sheet, row1based=i,
                        amount=_money(amount), nights=_money(_decimal(_cell(row, cols, "间夜数"))),
                        business_month=f"{period[:4]}-{period[4:6]}", countable=False,
                        date_basis="checkout_date")
                    controls.append(control)
                    result["summaries"].append(control)
            continue
        handled.add(sheet)
        h, cols = found
        for index, row in enumerate(rows[h + 1:], h + 2):
            if not any(_text(v) for v in row):
                continue
            if _text(row[0]).startswith("实际付款金额"):
                match = re.search(r"[:：]\s*([-\d,.]+)", _text(row[0]))
                if match:
                    control = dict(kind="ctrip_footer_total", sheet=sheet, row1based=index,
                                   amount=_money(_decimal(match[1])), countable=False)
                    controls.append(control)
                    result["summaries"].append(control)
                continue
            order_id = _identifier(_cell(row, cols, "订单号"), result, sheet, index)
            amount = _decimal(_cell(row, cols, "结算价"))
            day = _day(_cell(row, cols, "离店日期"), datemode)
            if not order_id or amount is None:
                _issue(result, "ctrip_unparsed_row", "明细行订单号或结算金额无法识别，未计入合计", sheet, index, "error")
                continue
            raw_count += 1
            if day is None:
                _issue(result, "ctrip_checkout_date_unknown", "离店日期无法识别，不能确定业务月份", sheet, index, "error")
            raw_nights = _decimal(_cell(row, cols, "间夜"))
            nights += raw_nights or Decimal(0)
            evidence = dict(sheet=sheet, row1based=index, amount=_money(amount),
                            date=day, row_type=_text(_cell(row, cols, "订单类型")),
                            nights=_money(raw_nights))
            for key, label in money_fields.items():
                evidence[key] = _money(_decimal(_cell(row, cols, label)))
            if order_id not in orders:
                orders[order_id] = _fact(digest, sheet, index, "platform_order", Decimal(0), day,
                    suffix=order_id, category="ota", channel="ctrip", platform_order_id=order_id,
                    evidence_refs=[], fields={"net": "0.00", "date_basis": "checkout_date",
                    "cash_received": "unconfirmed", "checkin_date": _day(_cell(row, cols, "入住日期"), datemode)})
                orders[order_id]["description"] = "携程订单净结算（含有符号调整）"
            fact = orders[order_id]
            if fact["date"] != day:
                _issue(result, "ctrip_conflicting_checkout_dates", "同一订单调整行离店日期不同，月份需核实", sheet, index)
                fact["business_month"] = None
            fact["evidence_refs"].append(evidence)
            fact["amount"] = _money(Decimal(fact["amount"]) + amount)
            fact["fields"]["net"] = fact["amount"]
            for key in money_fields:
                if evidence[key] is not None:
                    fact["fields"][key] = _money(Decimal(fact["fields"].get(key, "0.00")) + Decimal(evidence[key]))
    result["facts"].extend(orders.values())
    total = sum((Decimal(f["amount"]) for f in orders.values()), Decimal(0))
    result["summaries"].append(dict(kind="ctrip_detail_total", amount=_money(total),
        raw_row_count=raw_count, order_count=len(orders), nights=_money(nights), countable=False))
    if not controls:
        _issue(result, "missing_independent_total", "携程明细缺少可识别的独立汇总金额")
    for control in controls:
        control["matches_detail"] = control["amount"] is not None and Decimal(control["amount"]) == total
        if not control["matches_detail"]:
            _issue(result, "ctrip_total_mismatch", "携程明细净额与汇总金额不符", control["sheet"], control["row1based"], "error")
        if control.get("nights") is not None and Decimal(control["nights"]) != nights:
            _issue(result, "ctrip_nights_mismatch", "携程间夜与汇总不符", control["sheet"], control["row1based"], "error")
    _issue(result, "ctrip_cash_unconfirmed", "携程账单按离店日归属，净结算额不代表已收到银行款项")
    return handled


def _douyin(result, sheets, digest, datemode):
    handled = set()
    level_totals = defaultdict(Decimal)
    counts = Counter()
    for sheet, rows in sheets.items():
        detail = _header(rows, {"订单编号", "商家应得", "订单实收金额"})
        summary = _header(rows, {"账单名称", "到账时间", "到账金额", "商品实收", "商家支出"})
        if not detail and not summary:
            continue
        handled.add(sheet)
        h, cols = detail or summary
        level = "detail" if detail else "daily" if sheet == "日账单" else "monthly"
        for index, row in enumerate(rows[h + 1:], h + 2):
            if not any(_text(v) for v in row):
                continue
            amount = _decimal(_cell(row, cols, "商家应得" if detail else "到账金额"))
            if amount is None:
                _issue(result, "douyin_unparsed_amount", "抖音金额无法识别，未计入合计", sheet, index, "error")
                continue
            gross = _decimal(_cell(row, cols, "订单实收金额" if detail else "商品实收"))
            fee_names = ["软件服务费", "平台撮合服务费", "达人服务费", "服务商服务费", "团长服务费", "职人激励金", "分期免息手续费", "保险费用", "增量宝"]
            fees = {name: _money(_decimal(_cell(row, cols, name))) for name in fee_names if name in cols}
            cost = sum((Decimal(v) for v in fees.values() if v is not None), Decimal(0)) if detail else _decimal(_cell(row, cols, "商家支出"))
            day = _day(_cell(row, cols, "结算时间" if detail else "到账时间"), datemode)
            fields = dict(gross=_money(gross), net=_money(amount), signed_cost=_money(cost),
                fees=fees, subsidy=_money(_decimal(_cell(row, cols, "平台补贴金额" if detail else "平台补贴"))),
                user_paid=_money(_decimal(_cell(row, cols, "用户实付金额" if detail else "用户商品实付"))),
                date_basis="platform_settlement_date" if detail else "platform_payout_date",
                stay_month="unconfirmed", bank_receipt="unconfirmed")
            level_totals[level] += amount
            counts[level] += 1
            if gross is not None and cost is not None and gross + cost != amount:
                _issue(result, "douyin_amount_equation_mismatch", "商品实收加有符号费用与到账净额不符，需核对其他调整项", sheet, index, "error")
            if detail:
                if day is None:
                    _issue(result, "douyin_settlement_date_unknown", "结算日期无法识别，到账月份待核实", sheet, index)
                order_id = _identifier(_cell(row, cols, "订单编号"), result, sheet, index)
                if not order_id:
                    _issue(result, "douyin_missing_order_id", "抖音明细订单号无法准确识别", sheet, index, "error")
                fact = _fact(digest, sheet, index, "platform_payout", amount, day,
                    category="ota", channel="douyin", platform_order_id=order_id, fields=fields)
                fact["description"] = "抖音结算明细（到账口径，入住月份待匹配）"
                # A payout month is not a stay/revenue recognition month.
                fact["business_month"] = None
                fields["cash_month"] = day[:7] if day else None
                fields["verification_date"] = _day(_cell(row, cols, "核销时间"), datemode)
                result["facts"].append(fact)
            else:
                result["summaries"].append(dict(kind=f"douyin_{level}_payout", sheet=sheet,
                    row1based=index, date=day, amount=_money(amount), fields=fields, countable=False))
    for level, total in level_totals.items():
        result["summaries"].append(dict(kind=f"douyin_{level}_total", amount=_money(total),
            count=counts[level], countable=False))
    if not counts["detail"]:
        _issue(result, "douyin_missing_detail", "抖音缺少可识别的订单结算明细，汇总不生成订单事实", severity="error")
    for level in ("daily", "monthly"):
        if level not in level_totals:
            _issue(result, "douyin_missing_control", f"抖音缺少{level}汇总核验")
        elif level_totals[level] != level_totals["detail"]:
            _issue(result, "douyin_level_total_mismatch", f"抖音{level}汇总与明细净额不符", severity="error")
    _issue(result, "douyin_cash_basis", "月表、日表及订单结算明细为同组资金的不同层级；仅明细生成事实，入住月份及银行实收待核对")
    return handled


def _cleaning(result, data, filename, sheets, digest):
    from app.services.monthly_close.cleaning_work_log import parse_cleaning_work_log
    from app.services.monthly_close.service_statement import ServiceStatementError
    handled = {s for s, rows in sheets.items() if _header(rows, {"日期", "正常打扫房间号", "续住房间"})}
    try:
        entries = parse_cleaning_work_log(data, filename, None)
    except ServiceStatementError as exc:
        _issue(result, "cleaning_parse_error", f"{exc}。请核对这一行的日期和房号；本次未导入打扫记录。", severity="error")
        return handled
    for ordinal, entry in enumerate(entries or []):
        fact = _fact(digest, entry.source_sheet, entry.source_row, "cleaning_work",
            day=entry.service_date, suffix=str(ordinal), category="cleaning",
            fields={"service_type": entry.service_type, "quantity": 1,
                    "date_basis": "service_date", "charge": "not_derived"})
        fact.update(room_ref=entry.room_ref, description="续住打扫" if entry.service_type == "instay_cleaning" else "正常打扫")
        result["facts"].append(fact)
    result["summaries"].append(dict(kind="cleaning_work_count", count=len(result["facts"]),
        service_counts=dict(Counter(f["fields"]["service_type"] for f in result["facts"])), countable=False))
    _issue(result, "cleaning_evidence_only", "保洁日期、房间及类型仅为工作证据；超出数量及表列次数不推导收费，也不写入工作记录")
    return handled


def parse_financial_workbook(data: bytes, filename: str) -> dict | None:
    """Recognize supported source layouts, returning reviewable facts and issues.

    None means no supported layout (caller should report unsupported_template).
    Recognized files always disclose unhandled nonempty worksheets. All numbers
    in monetary fields are exact decimal strings; source byte hashes bind keys.
    """
    sheets, datemode, repaired = _load(data, filename)
    families = set()
    for sheet, rows in sheets.items():
        if _header(rows, {"房间名称", "收款日期", "费用科目", "已收金额"}):
            families.add("utility")
        if _header(rows, {"订单号", "结算价", "离店日期", "间夜"}):
            families.add("ctrip")
        if _header(rows, {"订单编号", "商家应得", "订单实收金额"}) or _header(rows, {"账单名称", "到账时间", "到账金额", "商品实收", "商家支出"}):
            families.add("douyin")
        if _header(rows, {"日期", "正常打扫房间号", "续住房间"}):
            families.add("cleaning")
    if not families:
        return None
    family = sorted(families)[0]
    result = dict(kind="ota" if family in {"ctrip", "douyin"} else family,
                  facts=[], summaries=[], issues=[], templates=[])
    if len(families) > 1:
        _issue(result, "ambiguous_workbook", "同一文件含多种财务模板，需拆分或明确映射；未自动生成事实", severity="error")
        return result
    digest = sha256(data).hexdigest()
    for sheet in repaired:
        _issue(result, "worksheet_dimensions_recovered", "导出表维度元数据不完整，已只读恢复实际行列", sheet, severity="info")
    if family == "utility":
        handled = _utility(result, sheets, digest, datemode)
    elif family == "ctrip":
        handled = _ctrip(result, sheets, digest, datemode)
    elif family == "douyin":
        handled = _douyin(result, sheets, digest, datemode)
    else:
        handled = _cleaning(result, data, filename, sheets, digest)
    for sheet, rows in sheets.items():
        if sheet not in handled and any(_text(v) for row in rows for v in row):
            _issue(result, "unsupported_sheet", "非空工作表未识别，未计入事实或合计", sheet, severity="error")
    if not result["facts"]:
        _issue(result, "no_source_facts", "未提取到可核对的明细事实", severity="error")
    return result
