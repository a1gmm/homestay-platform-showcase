"""Read-only accountant delivery verification; never approves or closes a month."""
from __future__ import annotations

import hashlib
from datetime import date
from io import BytesIO
import json
from typing import Any
from urllib.parse import quote
from decimal import Decimal

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.monthly_close import (
    MONTHLY_CLOSE_STEP_KEYS, MonthlyCloseCycle, MonthlyCloseDocument, MonthlyCloseInboxItem,
)
from app.models.user import User, UserRole
from app.services.monthly_close.evidence import build_step_evidences, _SOURCE_LABELS
from app.services.monthly_close.finalization import _finalization_live_snapshot
from app.services.monthly_close.permissions import MonthlyClosePermissionDenied
from app.services.monthly_close.projection import build_monthly_close_projection

STEP_LABELS = {
    "source_collection": "月结原始资料", "order_integrity": "订单及房段完整性",
    "service_fees": "保洁和布草费用", "utilities": "水电和运营支出",
    "ota_statements": "平台账单核对", "exception_clearance": "月结异常处理",
    "preflight": "结算体检", "settlement_review": "业主结算复核",
}
DELIVERY_STEPS = MONTHLY_CLOSE_STEP_KEYS[:-1]


def _hash(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    separators=(",", ":"), default=str).encode()).hexdigest()


def _money(value: Any) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


from app.services.monthly_close.issue_locators import issue_order_locators as _issue_order_locators


async def _package_fingerprint(response, expected_checks: list[dict], expected_totals) -> str:
    """Exercise XLSX serialization and validate the generated room reconciliation."""
    body = b"".join([chunk async for chunk in response.body_iterator])
    workbook = openpyxl.load_workbook(BytesIO(body), read_only=True, data_only=False)
    try:
        if not {"结算说明", "分成明细", "收入明细", "房号核对"} <= set(workbook.sheetnames):
            raise ValueError("package sheets incomplete")
        rows = list(workbook["房号核对"].iter_rows(values_only=True))
        details = [row for row in rows[1:] if isinstance(row[0], int)]
        if len(details) != len(expected_checks):
            raise ValueError("package room coverage differs")
        for row, expected in zip(details, expected_checks):
            if (str(row[1]) != str(expected["room_name"])
                    or int(row[2]) != int(expected["order_count"])
                    or _money(row[3]) != _money(expected["snapshot"])
                    or _money(row[4]) != _money(expected["current"])
                    or _money(row[5]) != 0):
                raise ValueError("package room totals differ")
        total_rows = [row for row in rows[1:] if row[0] == "合计"]
        if len(total_rows) != 1 or any(
            _money(total_rows[0][column]) != _money(sum(
                (Decimal(str(r[field])) for r in expected_checks), Decimal(0)))
            for column, field in ((3, "snapshot"), (4, "current"), (5, "diff"))
        ):
            raise ValueError("package reconciliation aggregate differs")
        statement = list(workbook["分成明细"].iter_rows(values_only=True))
        header, values = statement[1], statement[2]
        for label, amount in (("房费", expected_totals.total_net_revenue),
                              ("业主总支出", -expected_totals.deducted_expenses),
                              ("净利润", expected_totals.actual_owner_amount)):
            if _money(values[header.index(label)]) != _money(amount):
                raise ValueError("package statement totals differ")
        # Hash cells rather than ZIP bytes: workbook creation timestamps vary.
        return _hash({sheet.title: list(sheet.iter_rows(values_only=True))
                      for sheet in workbook.worksheets})
    finally:
        workbook.close()


async def verify_delivery(db: AsyncSession, cycle: MonthlyCloseCycle,
                          actor: User | dict[str, Any]) -> dict[str, Any]:
    """Return fresh, fingerprinted delivery facts for exactly ``cycle.billing_month``.

    ``ready`` describes accountant deliverables, not owner consent, payment, or
    formal month closure. Questions are deduplicated source/step obligations.
    Only current active administrators can run the complete financial verifier.
    No writes, confirmations, audit mutations, or commits are performed.
    """
    actor_id = (actor.get("user_id") if isinstance(actor, dict)
                else getattr(actor, "user_id", None))
    with db.no_autoflush:
        current_actor = await db.scalar(select(User).where(User.user_id == actor_id)
                                       .execution_options(populate_existing=True))
        if current_actor is None or not current_actor.is_active or current_actor.role != UserRole.admin:
            raise MonthlyClosePermissionDenied("delivery_forbidden", "仅管理员可核验完整月结交付")
        current_cycle = await db.scalar(select(MonthlyCloseCycle).where(
            MonthlyCloseCycle.cycle_id == cycle.cycle_id).execution_options(populate_existing=True))
        if current_cycle is None:
            raise ValueError("月结周期不存在")
        try:
            return await _verify(db, current_cycle, current_actor)
        except Exception:
            # A late storage/export failure must fail closed without exposing
            # database messages, original guest details, or credentials.
            message = "本月交付证据暂时无法完整读取，请稍后重新核验；持续失败请联系管理员。"
            return {
                "status": "blocked",
                "evidence_hash": _hash({"cycle_id": current_cycle.cycle_id,
                                        "billing_month": current_cycle.billing_month,
                                        "error": "delivery_verification_unavailable"}),
                "checks": [{"key": "live_evidence", "label": "读取当前月结证据",
                            "status": "blocked", "detail": message}],
                "questions": [{"key": "live_evidence", "subject": "当前月结证据", "message": message}],
                "artifacts": [], "summary": message,
            }


async def _verify(db, cycle, actor) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    questions: list[dict[str, Any]] = []
    artifacts: list[dict[str, Any]] = []
    fingerprints: dict[str, Any] = {"version": 1, "cycle_id": cycle.cycle_id,
                                    "billing_month": cycle.billing_month}
    question_keys: set[str] = set()

    def check(key, label, ok, detail, *, informational=False):
        checks.append({"key": key, "label": label,
                       "status": "passed" if ok else "pending" if informational else "blocked",
                       "detail": detail})

    def ask(key, subject, message, **context):
        if key not in question_keys:
            question_keys.add(key)
            questions.append({"key": key, "subject": subject, "message": message, **context})

    def result():
        blocked = any(c["status"] == "blocked" for c in checks)
        summary = (f"{cycle.billing_month} 会计交付仍有待处理事项，请按下列问题补齐后重新核验。"
                   if blocked else f"{cycle.billing_month} 会计交付资料与业主结算包已核验，可交付复核。")
        summary += " 业主确认、打款登记和正式关账分别以业务记录为准。"
        return {"status": "blocked" if blocked else "ready",
                "evidence_hash": _hash({**fingerprints, "checks": checks, "questions": questions,
                                        "artifacts": artifacts}),
                "checks": checks, "questions": questions, "artifacts": artifacts,
                "summary": summary}

    try:
        evidences = await build_step_evidences(db, cycle)

        async def evidence_loader(_db, _cycle):
            return evidences

        snapshot = await _finalization_live_snapshot(db, cycle, evidence_loader=evidence_loader)
        projection = await build_monthly_close_projection(db, cycle, actor)
        documents = (await db.execute(select(
            MonthlyCloseDocument.document_id, MonthlyCloseDocument.source_type,
            MonthlyCloseDocument.filename, MonthlyCloseDocument.sha256,
            MonthlyCloseDocument.byte_size, MonthlyCloseDocument.content,
        ).where(MonthlyCloseDocument.cycle_id == cycle.cycle_id,
                MonthlyCloseDocument.is_active.is_(True))
          .order_by(MonthlyCloseDocument.document_id))).all()
    except Exception:
        check("live_evidence", "读取当前月结证据", False, "证据读取未完成，请重新核验；持续失败请联系管理员检查数据服务。")
        ask("live_evidence", "当前月结证据", "重新读取本月原件、账本和结算记录后再核验交付。")
        return result()

    from app.models.financial_case import FinancialCaseSource
    from app.services.financial_case.service import load_sources
    financial_sources = await load_sources(db, cycle.cycle_id)
    financial_content = dict((await db.execute(select(
        FinancialCaseSource.source_id, FinancialCaseSource.content).where(
        FinancialCaseSource.source_id.in_([s.source_id for s in financial_sources])))).all())
    fingerprints["financial_originals"] = []
    for source in financial_sources:
        content = financial_content.get(source.source_id) or b""
        actual_hash = hashlib.sha256(content).hexdigest()
        fingerprints["financial_originals"].append({"source_id": source.source_id,
            "actual_hash": actual_hash, "version": source.version,
            "parsed_hash": _hash(source.parsed), "decisions_hash": _hash(source.decisions)})
        intact = bool(content) and actual_hash == source.sha256
        check(f"original:{source.source_id}", source.filename, intact,
              "财务原件内容与保存时一致。" if intact else "财务原件内容与保存时不一致。")
        if not intact:
            ask(f"original:{source.source_id}", source.filename,
                "原件内容与保存时不一致，请重新上传核对后再生成费用方案。",
                source_id=source.source_id, step_key="source_collection")
    fingerprints["snapshot"] = snapshot
    fingerprints["steps"] = [e.evidence_hash for e in evidences]
    fingerprints["sources"] = [s.to_dict() for s in projection.sources]
    fingerprints["originals"] = [{"document_id": row.document_id,
                                  "actual_hash": hashlib.sha256(row.content or b"").hexdigest()}
                                 for row in documents]
    work_log_checks = {}
    for source in projection.sources:
        if (source.source_type == "cleaning_statement" and source.state == "needs_action"
                and source.documents
                and all(d.analysis_state in {"ready", "work_log_ready"} for d in source.documents)
                and any(d.analysis_state == "work_log_ready" for d in source.documents)):
            # Projection keeps recognized work logs actionable for navigation.
            # Delivery uses the live activity and fee reconciliation instead.
            from app.services.monthly_close.cleaning_work_log import parse_cleaning_work_log, compare_cleaning_work_log
            from app.services.monthly_close.work_record_fees import plan_work_fees
            pending = await db.scalar(select(MonthlyCloseInboxItem.item_id).where(
                MonthlyCloseInboxItem.cycle_id == cycle.cycle_id,
                MonthlyCloseInboxItem.source_type == "cleaning_statement",
                MonthlyCloseInboxItem.status.not_in(("confirmed", "dismissed"))).limit(1))
            for projected in source.documents:
                if projected.analysis_state != "work_log_ready":
                    continue
                original = next((d for d in documents if d.document_id == projected.document_id), None)
                if original is None:
                    continue
                entries = parse_cleaning_work_log(original.content, original.filename, cycle.billing_month)
                if entries is None:
                    continue
                comparison = await compare_cleaning_work_log(db, entries, cycle.billing_month, original.document_id)
                fees = await plan_work_fees(db, original, cycle.billing_month)
                work_log_checks[original.document_id] = {
                    "complete": not pending and bool(entries) and not comparison.differences
                                and not fees["actions"] and not fees["unresolved"],
                    "comparison": comparison.model_dump(mode="json"), "fees": fees,
                }
    fingerprints["work_log_reconciliation"] = work_log_checks
    bad_sources: set[str] = set()
    for source in projection.sources:
        key = f"source:{source.source_type}"
        label = _SOURCE_LABELS.get(source.source_type, "月结资料")
        originals = [d for d in documents if d.source_type == source.source_type]
        damaged = [d for d in originals if not d.content or len(d.content) != d.byte_size
                   or hashlib.sha256(d.content).hexdigest() != d.sha256]
        reconciled_work_logs = (source.source_type == "cleaning_statement" and source.state == "needs_action"
            and any(d.analysis_state == "work_log_ready" for d in source.documents)
            and all(d.analysis_state == "ready" or work_log_checks.get(d.document_id, {}).get("complete")
                    for d in source.documents))
        ok = (source.state in {"completed", "not_applicable"} or reconciled_work_logs) and not damaged
        if source.state == "not_applicable" and not source.not_applicable_reason:
            ok = False
        message = ("原件内容与保存记录不一致，请重新上传并核对。" if damaged else
                   "请上传本月原始文件并完成识别、映射和核对；本月无此项时请说明原因。")
        check(key, label, ok, "本月资料已处理或已有不适用说明。" if ok else message)
        if not ok:
            bad_sources.add(source.source_type)
            ask(key, label, message, source_id=source.source_id, step_key="source_collection")
    if not projection.sources:
        check("sources_missing", "月结资料要求", False, "本月资料要求尚未建立。")
        ask("sources_missing", "月结资料", "先建立本月资料清单，再上传原件或说明不适用原因。", step_key="source_collection")

    evidence_by_key = {e.step_key: e for e in evidences}
    order_locators = await _issue_order_locators(db, cycle.billing_month, evidences)
    for step_key in DELIVERY_STEPS:
        evidence = evidence_by_key.get(step_key)
        label = STEP_LABELS[step_key]
        ok = evidence is not None and not evidence.blocking_count
        confirmed = (evidence is not None
                     and snapshot["step_confirmation_hashes"].get(step_key) == evidence.evidence_hash)
        check(f"step:{step_key}", label, ok,
              "当前证据核对通过。" if ok else "仍有业务差异，请先处理后重新核验。")
        if not ok:
            issues = evidence.snapshot.get("issues", []) if evidence else []
            for issue in issues:
                if issue.get("source_type") in bad_sources:
                    continue
                code = str(issue.get("code") or "evidence_blocked")
                subject = issue.get("subject") or label
                resource = issue.get("resource_id") or issue.get("order_id") or issue.get("room_id") or subject
                locator = order_locators.get(issue.get("order_id"))
                if locator:
                    subject = locator
                ask(f"issue:{code}:{resource}", subject,
                    f"{locator + '：' if locator else ''}{issue.get('cause') or issue.get('message') or '业务证据尚未核对完成。'} {issue.get('next_step') or '请核对对应原件和业务记录。'}",
                    step_key=step_key, code=code)
            if not issues and not (step_key == "source_collection" and bad_sources):
                ask(f"step:{step_key}", label, "补齐本步骤业务证据并处理差异后，重新核验。", step_key=step_key)
        # A blocked fact is one obligation; don't also ask to confirm it yet.
        check(f"confirmation:{step_key}", f"{label}确认", confirmed,
              "确认与当前证据一致。" if confirmed else "当前证据尚未确认或原确认已过期。")
        if ok and not confirmed:
            ask(f"confirmation:{step_key}", label, "请复核本步骤当前结果并完成确认。", step_key=step_key)

    unsettled = bool(snapshot["unresolved_issues"] or snapshot["open_remediation_ids"]
                     or snapshot["unsafe_attempt_ids"])
    check("execution_clearance", "异常与执行结果", not unsettled,
          "无未完成异常或执行补救。" if not unsettled else "仍有异常、待验证执行或补救事项。")
    if unsettled and not any(e.blocking_count for e in evidences if e.step_key != "owner_confirmation"):
        ask("execution_clearance", "执行结果核验", "请核对未解决异常和待验证执行，完成补救后重新核验。", step_key="exception_clearance")
    check("preflight", "最终结算体检", not snapshot["preflight"]["blocking"],
          "结算体检通过。" if not snapshot["preflight"]["blocking"] else "结算体检仍有阻断项。")

    settlements = snapshot["settlements"]
    check("settlements_present", "本月业主结算单", bool(settlements),
          f"找到本月 {len(settlements)} 份结算单。" if settlements else "本月没有结算单，不能认定交付完成。")
    if not settlements and not any(q.get("code") == "settlements_missing" for q in questions):
        ask("settlements_present", "本月业主结算", "完成本月账务核对后，生成并复核业主结算单。", step_key="settlement_review")
    # A single valid owner's package cannot establish coverage of the whole month.
    from app.models.order import Order, OrderStatus, BookingType, Channel
    from app.models.order_room import OrderRoom
    from app.models.room import Room
    year, month = map(int, cycle.billing_month.split("-"))
    start = date(year, month, 1)
    end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    expected_rooms = (await db.execute(select(Room.room_id, Room.room_name, Room.owner_id)
        .join(OrderRoom, OrderRoom.room_id == Room.room_id)
        .join(Order, Order.order_id == OrderRoom.order_id)
        .where(OrderRoom.check_out_date >= start, OrderRoom.check_out_date < end,
               Order.is_deleted.is_(False), Order.order_status != OrderStatus.cancelled,
               Order.booking_type != BookingType.owner_self, Order.channel != Channel.self_used,
               Room.owner_id.is_not(None)).distinct().order_by(Room.room_id))).all()
    covered = {(entry["owner_id"], rid) for entry in settlements for rid in entry["room_ids"]}
    missing_rooms = [r for r in expected_rooms if (r.owner_id, r.room_id) not in covered]
    duplicate_owners = len({s["owner_id"] for s in settlements}) != len(settlements)
    coverage_ok = not missing_rooms and not duplicate_owners
    fingerprints["expected_rooms"] = [tuple(row) for row in expected_rooms]
    check("settlement_coverage", "本月房间结算覆盖", coverage_ok,
          "本月经营房间均有对应业主结算，未发现重复业主结算。" if coverage_ok
          else "本月仍有房间未纳入结算或业主结算重复。")
    if missing_rooms and settlements:
        ask("settlement_coverage", "本月房间结算覆盖",
            "请将以下本月房间纳入对应业主结算：" + "、".join(r.room_name for r in missing_rooms) + "。",
            step_key="settlement_review")
    if duplicate_owners:
        ask("duplicate_settlements", "重复业主结算", "请核对本月同一业主的重复结算记录，确保收入和费用不会重复交付。", step_key="settlement_review")
    fingerprints["packages"] = []
    for entry in settlements:
        sid = entry["settlement_id"]
        label = "本月业主结算包"
        try:
            from app.api.v1.export import (
                _settlement_export_context, _settlement_income_rows, export_settlement_package,
            )
            from app.services.owner_settlement import settlement_current_amounts
            settlement, owner, rooms, _ = await _settlement_export_context(db, sid)
            label = f"{owner.name if owner else '业主'} · {cycle.billing_month} 结算包"
            if settlement.billing_month != cycle.billing_month or owner is None or not settlement.items:
                raise ValueError("incomplete settlement context")
            income_rows, room_checks = await _settlement_income_rows(db, settlement)
            if not income_rows or not room_checks or any(_money(r["diff"]) for r in room_checks):
                raise ValueError("no reconciled financial evidence")
            room_ids = {i.room_id for i in settlement.items if i.room_id}
            if room_ids != {r["room_id"] for r in room_checks} or not room_ids <= rooms.keys():
                raise ValueError("room coverage differs")
            totals = settlement_current_amounts(settlement)
            for field, item_field in (("total_net_revenue", "net_revenue"),
                                      ("deducted_expenses", "owner_expenses"),
                                      ("actual_owner_amount", "owner_net_amount")):
                if _money(getattr(totals, field)) != _money(sum(
                        (Decimal(getattr(i, item_field)) for i in settlement.items), Decimal(0))):
                    raise ValueError("settlement header totals differ")
            pending_artifacts = []
            for internal in (False, True):
                response = await export_settlement_package(sid, db,
                    {"user_id": actor.user_id, "role": "admin"}, internal=internal)
                digest = await _package_fingerprint(response, room_checks, totals)
                kind = "internal_owner_package" if internal else "owner_package"
                fingerprints["packages"].append({"settlement_id": sid, "kind": kind, "hash": digest})
                pending_artifacts.append({"label": f"{label}（{'内部核账' if internal else '对外脱敏'}）",
                    "url": f"/api/v1/export/settlements/{quote(sid, safe='')}/package?internal={'true' if internal else 'false'}",
                    "kind": kind})
            artifacts.extend(pending_artifacts)
            check(f"package:{sid}", label, True, f"已实际生成内外结算包，{len(room_checks)} 间房收入和结算合计一致。")
        except Exception:
            check(f"package:{sid}", label, False, "结算包未通过原始收入、逐房合计或文件生成核验。")
            ask(f"package:{sid}", label, "请核对本月收入明细、费用和结算快照；修正差异后重新生成待确认结算单。已确认结算请使用正式更正流程。", step_key="settlement_review")

    statuses = [s["status"] for s in settlements]
    approved = bool(statuses) and all(s in {"confirmed", "paid"} for s in statuses)
    paid = bool(statuses) and all(s == "paid" for s in statuses)
    check("owner_approval", "业主正式确认", approved,
          "业主结算均已确认。" if approved else "仍待业主正式确认；会计交付不替代业主同意。", informational=True)
    check("payment", "打款登记", paid,
          "全部结算已有打款登记；更正差额仍以支付凭证核对。" if paid else "尚未全部登记打款，不代表款项已经支付。", informational=True)
    check("formal_close", "正式月结状态", cycle.status == "completed",
          "周期记录为已完成。" if cycle.status == "completed" else "周期尚未正式关账。", informational=True)
    return result()
