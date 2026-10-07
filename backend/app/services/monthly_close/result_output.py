"""One typed, evidence-bound result for chat and downloadable business formats.

Renderers accept server-owned results only. No model HTML, formulas, URLs, file
paths or executable templates enter a renderer. Export is a snapshot, not approval.
"""
from __future__ import annotations

import csv
from decimal import Decimal
from html import escape
import io
import json
import re
from typing import Literal
import zipfile

from pydantic import BaseModel, ConfigDict, Field

Format = Literal["txt", "md", "csv", "json", "xlsx", "docx", "pdf", "html"]
FORMATS = ("txt", "md", "csv", "json", "xlsx", "docx", "pdf", "html")
LABELS = {"txt": "文本", "md": "Markdown", "csv": "CSV", "json": "JSON", "xlsx": "Excel", "docx": "Word", "pdf": "PDF", "html": "网页"}
MEDIA = {"txt": "text/plain; charset=utf-8", "md": "text/markdown; charset=utf-8", "csv": "text/csv; charset=utf-8", "json": "application/json", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "pdf": "application/pdf", "html": "text/html; charset=utf-8"}


class OutputRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    style: Literal["detail", "summary", "table", "checklist"] = "detail"
    formats: list[Format] = Field(default_factory=list, max_length=8)


class ResultSection(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(max_length=160)
    columns: list[str] = Field(max_length=12)
    rows: list[list[str | int | None]] = Field(default_factory=list, max_length=5000)
    note: str = Field(default="", max_length=1000)


class ResultDocument(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str
    billing_month: str
    queried_at: str
    summary: str
    sections: list[ResultSection] = Field(default_factory=list, max_length=20)
    caveats: list[str] = Field(default_factory=list, max_length=20)


class ResultOutput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    document: ResultDocument
    presentation: OutputRequest


class FormattedResultFacts(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: Literal["formatted_result"] = "formatted_result"
    billing_month: str
    request_text: str
    source_run_id: str
    document: ResultDocument
    presentation: OutputRequest


def unsupported_output_request(text: str) -> bool:
    return bool(re.search(r"导出|输出|下载|转成|整理成|生成", text)
                and re.search(r"(?<![a-zA-Z0-9_])(?:pptx?|powerpoint|xml|ya?ml|zip|png|jpe?g)(?![a-zA-Z0-9_])|图片文件", text, re.I))


def output_request(text: str) -> OutputRequest | None:
    formats = []
    for name, pattern in [("xlsx", r"excel|xlsx|电子表格"), ("docx", r"word|docx|文档版"), ("pdf", r"pdf|打印版"), ("csv", r"csv"), ("json", r"json"), ("md", r"markdown|\.md\b"), ("html", r"html|网页版"), ("txt", r"txt|纯文本|文本文件")]:
        if re.search(pattern, text, re.I):
            formats.append(name)
    style = "summary" if re.search(r"简短|简洁|简要|一句话|给老板|摘要|总结版", text) else "table" if re.search(r"表格|列表格|分列|列成表", text) else "checklist" if re.search(r"待办清单|核实清单|检查清单|逐项清单", text) else "detail"
    if not formats and style == "detail":
        return None
    # Asking what columns an incoming file needs is not asking for an export.
    if re.search(r"上传|需要什么格式|要哪些列|要什么格式|表头|识别失败", text) and not re.search(r"导出|下载|输出|生成.*(?:报告|文件)|整理成", text):
        return None
    return OutputRequest(style=style, formats=list(dict.fromkeys(formats)))


def formatting_only(text: str) -> bool:
    if not output_request(text) and not unsupported_output_request(text):
        return False
    value = re.sub(r"(?:不要|别)(?:修改|执行|记账|入账)(?:账目|方案)?", "", text)
    if re.search(r"删除|补齐|补录|执行|批准|确认|恢复|撤销|改成公司|修改金额|记账|入账|转账", value):
        return False
    # A new date, room, category or recalculation is a new query, not a reformat.
    if re.search(r"20\d{2}|\d+月|本月|这个月|整月|\d{3,4}房|重新核对|重新查询|最新|多少钱|多少笔|保洁|布草|银行|订单|水电|费用|工资|利润|具体|逐项|逐笔|为什么|原因|有哪些|第.{1,4}[项条笔]", value):
        return False
    return bool(re.search(r"这些|这份|这个|上述|刚才|上面|结果|改成|整理成|换成|转成|输出|导出|下载|给我|给老板|用.*(?:表|格式|清单)|简短|简洁", value))


def scope_counted_document(document: ResultDocument, text: str) -> ResultDocument:
    """A counted export must select one actual displayed group, never widen it."""
    match = re.search(r"[这那](\d+|[零一二两三四五六七八九十百]+)[项条笔]", text)
    if not match:
        return document
    from app.services.monthly_close.investigation_context import _number
    sections = [section for section in document.sections if len(section.rows) == _number(match[1])]
    if len(sections) != 1:
        raise LookupError("原结果中没有唯一对应这个数量的清单，请先展开要导出的那组事项，再指定文件格式。")
    return document.model_copy(update={"sections": sections})


def _cell(value):
    if value is None:
        return "待确认"
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def present_document(document: ResultDocument, request: OutputRequest, text: str) -> ResultDocument:
    if request.style != "summary":
        return document
    # Shorten server-owned prose without generating new amounts or conclusions.
    sentences = [value.strip() for value in re.split(r"[。！？\n]+", document.summary) if value.strip()]
    count = 1 if "一句话" in text else 2
    return document.model_copy(update={"summary": "。".join(sentences[:count]) + "。"})


def build_document(reply, queried_at: str) -> ResultDocument:
    facts = reply.facts
    if reply.tool == "formatted_result":
        return ResultDocument.model_validate(facts["document"])
    month = str(facts.get("billing_month", ""))
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", month):
        raise ValueError("本条回复没有可核验的月份范围，请重新查询后导出。")
    sections = []
    def section(title, columns, rows, note=""):
        if rows:
            sections.append(ResultSection(title=title, columns=columns, rows=[[_cell(value) for value in row] for row in rows], note=note))
    details = facts.get("detail_items") or []
    section("逐项核对", ["序号", "事项", "日期", "房间", "差异与原因", "依据", "下一步"],
            [[r.get("ordinal", i + 1), r.get("subject"), r.get("service_date", "—"), r.get("room", "—"), r.get("cause"), r.get("source", "见系统对应记录"), r.get("next_step")] for i, r in enumerate(details)])
    focused_details = facts.get("detail_total") is not None and facts.get("focus") not in {None, "all"}
    section("资料与待办", ["资料", "当前结果", "下一步"], [[r.get("label"), r.get("detail"), r.get("next_step")] for r in facts.get("sources", []) if "detail" in r and not focused_details])
    step_states = {"confirmed": "已完成流程确认", "blocked": "有事项待处理", "ready": "可进行流程复核", "locked": "等待前序步骤", "stale": "依据变化，需重新核对"}
    section("流程复核", ["步骤", "状态", "阻断项数"], [[r.get("label"), step_states.get(r.get("status"), "待核实"), r.get("blocking_count")] for r in facts.get("steps", []) if not focused_details])
    amount = facts.get("amount_summary") or {}
    section("已入账费用", ["分类", "笔数", "金额（元）"], [[r["label"], r["count"], r["amount"]] for r in amount.get("subtotals", [])], "按原查询范围与业务口径统计；不等于供应商已付款或最终利润。")
    comparison = facts.get("comparison") or {}
    section("打扫记录差异", ["日期", "房间", "类型", "原表次数", "系统条数", "原表行"],
            [[r.get("service_date"), r.get("room_ref"), "续住打扫" if r.get("service_type") == "instay_cleaning" else "正常打扫", r.get("table_count"), r.get("system_count"), "、".join(map(str, r.get("source_rows", [])))] for r in comparison.get("differences", [])])
    section("核对指标", ["项目", "结果"], [[r.get("label"), r.get("value")] for r in facts.get("metrics", [])])
    section("依据明细", ["项目", "依据"], [[r.get("label"), r.get("value")] for r in facts.get("details", [])])
    section("待核实事项", ["事项", "说明", "下一步"], [[r.get("label", r.get("subject", "待核实事项")), r.get("message", r.get("cause")), r.get("next_step", "核实原件与业务事实后继续")] for r in facts.get("issues", [])])
    cases = [facts["selected"]] if facts.get("selected") else facts.get("cases", [])
    section("费用调查", ["日期", "房间", "供应商金额（元）", "系统金额（元）", "差额（元）", "说明", "待核实"], [[r.get("service_date"), r.get("room_id"), r.get("vendor_amount"), r.get("system_amount"), r.get("difference"), r.get("summary"), r.get("question")] for r in cases])
    if reply.tool == "cleaning_work_chat":
        fee_summary = facts.get("fee_summary") or {}
        section("续住费用待核实", ["日期", "房间", "确认次数", "待核实原因", "下一步"],
                [[r.get("service_date"), r.get("room_ref"), r.get("quantity"), r.get("reason"), "核实原费用及收费依据后，再提出更正方案"] for r in fee_summary.get("unresolved", [])],
                "这些记录尚未补记费用；不能把零元、作废或无法明确对应的费用直接重新收取。")
        if facts.get("state") == "report":
            section("续住费用待补记候选", ["日期", "房间", "确认次数", "单价（元）", "已有金额（元）", "拟补金额（元）"],
                    [[r.get("service_date"), r.get("room_ref"), r.get("quantity"), r.get("unit_price"), r.get("existing_amount"), r.get("amount")] for r in fee_summary.get("actions", [])],
                    "仅查询候选，不是可执行方案；需要另行生成并确认方案，本次未修改费用。")
        section("续住费用已有处理", ["日期", "房间", "确认次数", "处理依据"],
                [[r.get("service_date"), r.get("room_ref"), r.get("quantity"), r.get("reason")] for r in fee_summary.get("recognized", [])],
                "保留已有费用结果，不自动补收；记账结果不等同实际付款。")
        kinds = {"add": "补记打扫记录", "delete": "撤下打扫记录", "restore": "恢复打扫记录", "resolve": "确认实际次数", "fee": "补记费用"}
        rows = []
        for action in facts.get("actions", []):
            detail = action.get("reason") or action.get("description") or "按已核实的原表与系统记录处理"
            if action.get("confirmed_count") is not None:
                detail += f"；确认次数：{action['confirmed_count']}"
            if action.get("related_fees_unchanged"):
                detail += "；关联费用保持原记录，不随打扫记录撤下"
            rows.append([kinds.get(action.get("kind"), "待核实动作"), action.get("service_date"), action.get("room_ref"), action.get("amount", "不涉及金额变更"), detail])
        section("处理动作", ["动作", "日期", "房间", "金额（元）", "处理说明"], rows,
                "待确认方案尚未执行。" if facts.get("state") == "proposal" else "动作状态以本次回复与系统复核结果为准。")
    if reply.tool == "order_identity_chat" and facts.get("order_id"):
        section("订单身份核对", ["系统订单", "平台订单号", "状态"], [[facts["order_id"], facts.get("platform_order_id"), {"proposal": "待确认", "completed": "已完成", "conflict": "记录变化，需重新核对"}.get(facts.get("state"), "待核实")]])
    caveats = ["本文件是本次核对的查询快照，不是结算凭证，也不表示已经改账、付款或完成关账。", "资料、账本或解释变化后需要重新查询；未知金额与未核实事实不能视为零。"]
    if facts.get("detail_total") is not None and len(details) < facts["detail_total"]:
        caveats.append(f"原查询共有 {facts['detail_total']} 项，本快照仅含已展示的 {len(details)} 项；这是原查询的展示范围；如需全部，请重新查询完整清单。")
    if reply.intent == "action_plan":
        caveats.insert(0, "这是待确认方案的说明，尚未执行；下载文件不会批准或执行方案。")
    return ResultDocument(title=f"{month} 对账核对结果", billing_month=month, queried_at=queried_at,
                          summary=reply.message, sections=sections, caveats=caveats)


def text_document(document: ResultDocument, *, markdown=False) -> str:
    lines = [("# " if markdown else "") + document.title, f"查询时间：{document.queried_at}", document.summary, *document.caveats]
    for section in document.sections:
        lines.extend(["", ("## " if markdown else "") + section.title])
        if section.note:
            lines.append(section.note)
        if markdown:
            clean = lambda value: escape(_cell(value)).replace("|", "\\|").replace("\n", " ")
            lines.append("| " + " | ".join(map(clean, section.columns)) + " |")
            lines.append("| " + " | ".join("---" for _ in section.columns) + " |")
            lines.extend("| " + " | ".join(map(clean, row)) + " |" for row in section.rows)
        else:
            lines.append("\t".join(section.columns))
            lines.extend("\t".join(_cell(value).replace("\t", " ") for value in row) for row in section.rows)
    return "\n".join(lines)


def _spreadsheet_text(value):
    text = _cell(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def render_document(document: ResultDocument, fmt: Format) -> bytes:
    if fmt == "txt":
        return text_document(document).encode("utf-8")
    if fmt == "md":
        return text_document(document, markdown=True).encode("utf-8")
    if fmt == "json":
        return document.model_dump_json(indent=2).encode("utf-8")
    if fmt == "csv":
        out = io.StringIO(newline="")
        writer = csv.writer(out)
        # Different sections have different schemas. A rectangular long-form
        # CSV can be read by Excel/pandas without malformed mixed-width rows.
        writer.writerow(["分区", "记录序号", "字段", "内容"])
        for key, value in [("标题", document.title), ("查询时间", document.queried_at), ("结论", document.summary), *[("口径说明", text) for text in document.caveats]]:
            writer.writerow(["查询说明", "", key, _spreadsheet_text(value)])
        for section in document.sections:
            if section.note:
                writer.writerow([section.title, "", "口径说明", _spreadsheet_text(section.note)])
            for index, row in enumerate(section.rows, 1):
                for column, value in zip(section.columns, row):
                    writer.writerow([section.title, index, column, _spreadsheet_text(value)])
        return out.getvalue().encode("utf-8-sig")
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        book = Workbook()
        first = book.active
        first.title = "先看这页"
        for line in [document.title, f"查询时间：{document.queried_at}", document.summary, *document.caveats]:
            first.append([_spreadsheet_text(line)])
        first.column_dimensions["A"].width = 100
        first.sheet_view.showGridLines = False
        for row in first:
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
            first.row_dimensions[row[0].row].height = 45 if len(str(row[0].value)) < 100 else 110
        for section in document.sections:
            sheet = book.create_sheet(section.title[:31])
            sheet.append(section.columns)
            for row in section.rows:
                values = []
                for col, value in zip(section.columns, row):
                    if re.fullmatch(r"(?:.*金额（元）|差额（元）|笔数|序号|阻断项数|原表次数|系统条数)", col) and re.fullmatch(r"-?\d+(?:\.\d+)?", str(value)) and len(str(value).replace('.', '').lstrip('-0')) <= 15:
                        values.append(Decimal(str(value)))
                    else:
                        values.append(_spreadsheet_text(value))
                sheet.append(values)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            for index in range(1, len(section.columns) + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 22 if index < 3 else 38
            for row in sheet:
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if cell.row == 1:
                        cell.fill = PatternFill("solid", fgColor="E5DDCB")
                        cell.font = Font(name="Calibri", bold=True)
            sheet.print_title_rows = "1:1"
        out = io.BytesIO(); book.save(out); return out.getvalue()
    if fmt == "html":
        body = "<h1>" + escape(document.title) + "</h1><p>查询时间：" + escape(document.queried_at) + "</p><p>" + escape(document.summary).replace("\n", "<br>") + "</p>"
        body += "<ul>" + "".join("<li>" + escape(line) + "</li>" for line in document.caveats) + "</ul>"
        for section in document.sections:
            body += "<h2>" + escape(section.title) + "</h2><p>" + escape(section.note) + "</p><div class='table'><table><thead><tr>" + "".join("<th>" + escape(col) + "</th>" for col in section.columns) + "</tr></thead><tbody>"
            body += "".join("<tr>" + "".join("<td>" + escape(_cell(value)) + "</td>" for value in row) + "</tr>" for row in section.rows) + "</tbody></table></div>"
        return ("<!doctype html><html lang='zh-CN'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; style-src 'unsafe-inline'\"><title>对账核对结果</title><style>body{font:16px/1.7 system-ui;margin:24px;color:#2b2721;background:#fbf8f1}h1,h2{font-weight:500}.table{overflow:auto}table{border-collapse:collapse;width:100%}th,td{border:1px solid #e5ddcb;padding:10px;text-align:left;vertical-align:top;min-width:90px}th{background:#f5f1ea}</style>" + body + "</html>").encode("utf-8")
    if fmt == "docx":
        def para(value, heading=False):
            content = escape(_cell(value)).replace("\n", '</w:t><w:br/><w:t xml:space="preserve">')
            keep = '<w:keepNext/>' if heading else ''
            weight = '<w:b/><w:sz w:val="26"/>' if heading else '<w:sz w:val="22"/>'
            return '<w:p><w:pPr>' + keep + '<w:spacing w:after="120"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Arial" w:eastAsia="Microsoft YaHei"/>' + weight + '</w:rPr><w:t xml:space="preserve">' + content + '</w:t></w:r></w:p>'
        body = para(document.title, True) + "".join(para(line) for line in ["查询时间：" + document.queried_at, document.summary, *document.caveats])
        for section in document.sections:
            body += para(section.title, True) + (para(section.note) if section.note else '')
            # Per-item labeled paragraphs remain readable for wide business evidence.
            for index, row in enumerate(section.rows, 1):
                body += para(f"{index}.") + "".join(para(f"{col}：{_cell(value)}") for col, value in zip(section.columns, row))
        xml = '<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + body + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1000" w:right="1000" w:bottom="1000" w:left="1000"/></w:sectPr></w:body></w:document>'
        files = {'[Content_Types].xml': '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>', '_rels/.rels': '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>', 'word/document.xml': xml}
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in files.items(): archive.writestr(name, content)
        return out.getvalue()
    if fmt == "pdf":
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, LongTable, TableStyle
        font = "STSong-Light"
        if font not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(UnicodeCIDFont(font))
        style = ParagraphStyle("Business", fontName=font, fontSize=10, leading=16, wordWrap="CJK", spaceAfter=7)
        title = ParagraphStyle("Title", parent=style, fontSize=18, leading=25, spaceAfter=16)
        cell_style = ParagraphStyle("Cell", parent=style, fontSize=9, leading=13, spaceAfter=0)
        page_size = landscape(A4) if any(len(s.columns) >= 6 for s in document.sections) else A4
        available_width = page_size[0] - 84
        def width_weight(column):
            if re.search("原因|依据|说明|下一步|来源", column): return 3
            if re.search("序号|次数|条数|笔数|房间|状态", column): return 1
            return 1.5
        story = [Paragraph(escape(document.title), title)]
        for line in ["查询时间：" + document.queried_at, document.summary, *document.caveats]:
            story.append(Paragraph(escape(line).replace("\n", "<br/>"), style))
        for section in document.sections:
            story.extend([Spacer(1, 12), Paragraph(escape(section.title), title)])
            if section.note: story.append(Paragraph(escape(section.note), style))
            if section.columns and section.rows:
                weights = [width_weight(col) for col in section.columns]
                # The bundled Chinese CID font lacks the middle-dot separator.
                rows = [[Paragraph(escape(_cell(value).replace("·", " / ")).replace("\n", "<br/>"), cell_style) for value in row]
                        for row in [section.columns, *section.rows]]
                table = LongTable(rows, colWidths=[available_width * w / sum(weights) for w in weights],
                                  repeatRows=1, splitByRow=1, splitInRow=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5DDCB")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#E5DDCB")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.append(table)
        out = io.BytesIO()
        def footer(canvas, doc):
            canvas.setFont(font, 9); canvas.drawRightString(page_size[0] - 42, 25, f"第 {doc.page} 页  查询快照")
        SimpleDocTemplate(out, pagesize=page_size, leftMargin=42, rightMargin=42, topMargin=42, bottomMargin=42, title=document.title).build(story, onFirstPage=footer, onLaterPages=footer)
        return out.getvalue()
    raise ValueError("不支持该格式")
